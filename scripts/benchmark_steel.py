from __future__ import annotations

import argparse
import json
import re
import uuid
from pathlib import Path

from openpyxl import load_workbook

from server.documents import parse_docx
from server.matcher import _rank_matches, build_scope_sample, infer_scope, match_clauses


def normalized(value: object) -> str:
    return re.sub(r"[^0-9a-z가-힣]", "", str(value or "").casefold())


def main() -> None:
    parser = argparse.ArgumentParser(description="Measure steel matching against a legacy reference set.")
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--reference", type=Path, required=True)
    parser.add_argument("--kcs-raw", type=Path, required=True)
    parser.add_argument("--full-audit", action="store_true")
    args = parser.parse_args()

    titlelish, clauses, _ = parse_docx(
        args.source.read_bytes(), args.source.name, str(uuid.uuid4())
    )
    sheet = load_workbook(args.reference, read_only=True, data_only=True)["매칭결과"]
    reference_rows = [row for row in sheet.iter_rows(min_row=2, values_only=True) if row[4] == 1]
    clause_index = {
        (normalized(clause["label"]), normalized(clause["title"])): clause
        for clause in clauses
        if clause.get("source_type") in {"paragraph", "table"}
    }
    benchmark: list[tuple[dict, str]] = []
    for row in reference_rows:
        clause = clause_index.get((normalized(row[1]), normalized(row[2])))
        if clause is not None:
            benchmark.append((dict(clause), str(row[5])))

    prefixes, display = infer_scope(
        args.source.name, titlelish, build_scope_sample(clauses), args.kcs_raw
    )
    selected = [clause for clause, _ in benchmark]
    match_clauses(selected, args.kcs_raw, prefixes)
    rank1 = 0
    top3 = 0
    no_candidate = 0
    misses: list[dict[str, object]] = []
    for clause, (_, expected) in zip(selected, benchmark):
        codes = [candidate["kcs_code"] for candidate in clause.get("candidates", [])]
        no_candidate += not codes
        rank1 += bool(codes and codes[0] == expected)
        top3 += expected in codes
        if expected not in codes:
            source_text = " ".join((
                str(clause.get("match_context") or ""),
                str(clause["title"]),
                str(clause["title"]),
                str(clause["content"]),
            ))
            raw_ranked = _rank_matches(source_text, args.kcs_raw, prefixes)
            expected_rows = [
                (rank, round(score, 4), section["title"])
                for rank, (section, score) in enumerate(raw_ranked, start=1)
                if section["code"] == expected
            ]
            misses.append({
                "label": clause["label"],
                "title": clause["title"],
                "expected": expected,
                "actual": codes,
                "expected_retrieval": expected_rows[:1],
            })

    total = len(benchmark)
    report: dict[str, object] = {
        "reference_kind": "legacy automatic rank-1 (not human gold)",
        "reference_rows": len(reference_rows),
        "mapped_rows": total,
        "scope": display,
        "rank1_retained": rank1,
        "rank1_retention_rate": round(rank1 / total, 4) if total else None,
        "top3_retained": top3,
        "top3_retention_rate": round(top3 / total, 4) if total else None,
        "no_candidate": no_candidate,
        "misses": misses,
    }
    if args.full_audit:
        reviewable = [
            dict(clause) for clause in clauses
            if clause.get("source_type") in {"paragraph", "table"}
        ]
        match_clauses(reviewable, args.kcs_raw, prefixes)
        with_candidates = [clause for clause in reviewable if clause.get("candidates")]
        primary_top1 = sum(
            candidate["kcs_code"].replace(" ", "").startswith(("KCS4131", "KCS1431"))
            for clause in with_candidates
            for candidate in clause["candidates"][:1]
        )
        report["full_audit"] = {
            "reviewable_clauses": len(reviewable),
            "clauses_with_candidates": len(with_candidates),
            "candidate_coverage_rate": round(len(with_candidates) / len(reviewable), 4),
            "primary_steel_scope_top1": primary_top1,
            "primary_steel_scope_top1_rate": round(primary_top1 / len(with_candidates), 4)
            if with_candidates else None,
        }
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
