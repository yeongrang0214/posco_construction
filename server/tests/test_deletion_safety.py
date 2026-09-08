from __future__ import annotations

import io
import json
from datetime import datetime, timezone

import pytest
from docx import Document
from fastapi.testclient import TestClient

import server.app as app_module
from server.storage import Store


def _candidate(candidate_id: str, warnings: list[str] | None = None) -> dict:
    return {
        "id": candidate_id,
        "rank": 1 if candidate_id == "safe-candidate" else 2,
        "kcs_code": "KCS 41 31 05",
        "document_name": "건축물 강구조공사 일반사항",
        "version": "2026",
        "update_date": "2026-09-01",
        "kcs_clause": "3.2.1",
        "title": "강구조 요구사항",
        "content": "강구조 요구사항을 적용한다.",
        "score": 0.72,
        "classification": "관련성 높음",
        "reasons": ["공통 핵심어"],
        "warnings": warnings or [],
    }


def _store_with_project(tmp_path) -> tuple[Store, str]:
    project_id = "safety-project"
    store = Store(tmp_path / "safety.sqlite3")
    store.create_project(
        {
            "id": project_id,
            "title": "삭제 안전성 시험",
            "source_filename": "source.docx",
            "source_path": str(tmp_path / "source.docx"),
            "source_sha256": "hash",
            "uploaded_at": datetime.now(timezone.utc).isoformat(),
            "kcs_snapshot": "2026-09-04T22:34:56+09:00",
            "kcs_revision": "test-revision",
            "kcs_scope": "KCS 41 31",
            "warning": "",
        },
        [
            {
                "id": "heading",
                "source_order": 1,
                "label": "1",
                "title": "일반사항",
                "content": "일반사항",
                "source_type": "heading",
                "outline_level": 1,
                "candidates": [],
            },
            {
                "id": "clause",
                "source_order": 2,
                "label": "1.1",
                "title": "강구조 요구사항",
                "content": "강구조 요구사항을 적용한다.",
                "source_type": "paragraph",
                "outline_level": None,
                "candidates": [
                    _candidate("safe-candidate"),
                    _candidate("warning-candidate", ["수치 차이: 10 mm / 20 mm"]),
                ],
            },
        ],
    )
    return store, project_id


def test_decision_reason_heading_and_delete_gate(tmp_path):
    store, project_id = _store_with_project(tmp_path)

    with pytest.raises(ValueError, match="사유"):
        store.update_decision(project_id, "clause", "keep", "", "", None)
    with pytest.raises(ValueError, match="판정 대상"):
        store.update_decision(
            project_id,
            "heading",
            "keep",
            "",
            "",
            None,
            decision_reason="posco_specific",
        )
    with pytest.raises(ValueError, match="KCS 근거"):
        store.update_decision(
            project_id,
            "clause",
            "delete",
            "",
            "",
            None,
            decision_reason="fully_covered_by_kcs",
            coverage_confirmed=True,
        )
    with pytest.raises(ValueError, match="전체"):
        store.update_decision(
            project_id,
            "clause",
            "delete",
            "",
            "",
            "safe-candidate",
            decision_reason="fully_covered_by_kcs",
        )
    with pytest.raises(ValueError, match="차이 경고"):
        store.update_decision(
            project_id,
            "clause",
            "delete",
            "",
            "",
            "warning-candidate",
            decision_reason="fully_covered_by_kcs",
            coverage_confirmed=True,
        )

    saved = store.update_decision(
        project_id,
        "clause",
        "delete",
        "",
        "",
        "safe-candidate",
        decision_reason="fully_covered_by_kcs",
        coverage_confirmed=True,
    )
    assert saved and saved["coverage_confirmed"] is True
    project = store.get_project(project_id)
    assert project and project["reviewable_clauses"] == 1
    assert project["unreviewed_clauses"] == 0
    assert project["unsafe_delete_count"] == 0
    assert project["review_submission_ready"] is True


def test_non_kcs_delete_does_not_require_a_candidate(tmp_path):
    store, project_id = _store_with_project(tmp_path)

    saved = store.update_decision(
        project_id,
        "clause",
        "delete",
        "",
        "사내 시방서의 다른 조항과 중복됨",
        "safe-candidate",
        decision_reason="internal_duplicate",
        coverage_confirmed=True,
    )

    assert saved is not None
    assert saved["decision_reason"] == "internal_duplicate"
    assert saved["selected_candidate_id"] is None
    assert saved["coverage_confirmed"] is False
    project = store.get_project(project_id)
    assert project is not None
    assert project["unsafe_delete_count"] == 0
    assert project["invalid_decision_count"] == 0
    assert project["review_submission_ready"] is True


def test_no_kcs_match_keep_clears_an_existing_candidate(tmp_path):
    store, project_id = _store_with_project(tmp_path)

    saved = store.update_decision(
        project_id,
        "clause",
        "keep",
        "강구조 요구사항을 적용한다.",
        "",
        "safe-candidate",
        decision_reason="no_kcs_match",
        coverage_confirmed=True,
    )

    assert saved is not None
    assert saved["decision"] == "keep"
    assert saved["decision_reason"] == "no_kcs_match"
    assert saved["selected_candidate_id"] is None
    assert saved["coverage_confirmed"] is False


def test_store_startup_repairs_legacy_no_kcs_match_selection(tmp_path):
    store, project_id = _store_with_project(tmp_path)
    store.update_decision(
        project_id,
        "clause",
        "keep",
        "강구조 요구사항을 적용한다.",
        "",
        None,
        decision_reason="no_kcs_match",
    )
    with store.connect() as connection:
        connection.execute(
            """
            UPDATE clauses
            SET selected_candidate_id = 'safe-candidate', coverage_confirmed = 1
            WHERE id = 'clause'
            """
        )
        history_before = connection.execute(
            "SELECT COUNT(*) FROM decision_history WHERE clause_id = 'clause'"
        ).fetchone()[0]

    repaired_store = Store(store.database_path)
    repaired = repaired_store.get_clause(project_id, "clause")
    assert repaired is not None
    assert repaired["selected_candidate_id"] is None
    assert repaired["coverage_confirmed"] is False
    with repaired_store.connect() as connection:
        history_after = connection.execute(
            "SELECT COUNT(*) FROM decision_history WHERE clause_id = 'clause'"
        ).fetchone()[0]
    assert history_after == history_before + 1

    Store(store.database_path)
    with repaired_store.connect() as connection:
        history_after_second_start = connection.execute(
            "SELECT COUNT(*) FROM decision_history WHERE clause_id = 'clause'"
        ).fetchone()[0]
    assert history_after_second_start == history_after


def test_management_delete_requires_a_written_reason(tmp_path):
    store, project_id = _store_with_project(tmp_path)

    with pytest.raises(ValueError, match="검토의견"):
        store.update_decision(
            project_id,
            "clause",
            "delete",
            "",
            "",
            None,
            decision_reason="management_decision",
        )

    saved = store.update_decision(
        project_id,
        "clause",
        "delete",
        "",
        "현장 운영 범위에서 제외하기로 협의함",
        None,
        decision_reason="management_decision",
    )

    assert saved is not None
    assert saved["decision_reason"] == "management_decision"
    project = store.get_project(project_id)
    assert project is not None
    assert project["invalid_decision_count"] == 0
    assert project["review_submission_ready"] is True


def test_final_export_blocks_until_ready_and_removes_review_metadata(tmp_path, monkeypatch):
    store, project_id = _store_with_project(tmp_path)
    monkeypatch.setattr(app_module, "store", store)
    monkeypatch.setattr(
        app_module,
        "_reconcile_current_kcs_revision",
        lambda: ("2026-09-04T22:34:56+09:00", {"revision": "test-revision"}),
    )

    with TestClient(app_module.app) as client:
        blocked = client.get(f"/api/projects/{project_id}/export/final")
        assert blocked.status_code == 409
        assert "미검토 조항 1개" in blocked.json()["detail"]

        saved = client.patch(
            f"/api/projects/{project_id}/clauses/clause",
            json={
                "decision": "keep",
                "decision_reason": "posco_stricter",
                "coverage_confirmed": False,
                "edited_content": "포스코 강화 요구사항만 적용한다.",
                "review_note": "내부 검토 메모",
                "selected_candidate_id": "safe-candidate",
            },
        )
        assert saved.status_code == 200
        assert saved.json()["project"]["review_submission_ready"] is True
        assert saved.json()["project"]["final_export_ready"] is False

        submitted = client.post(
            f"/api/projects/{project_id}/review-workflow/submit",
            json={"author_name": "작성자", "note": "최종본 승인 요청"},
        )
        assert submitted.status_code == 200
        submission_id = submitted.json()["workflow"]["latest_review_submission"]["id"]
        approved = client.post(
            f"/api/projects/{project_id}/review-workflow/{submission_id}/approve",
            json={"approver_name": "승인자", "note": "최종본 승인"},
        )
        assert approved.status_code == 200
        assert approved.json()["project"]["final_export_ready"] is True

        final = client.get(f"/api/projects/{project_id}/export/final")
        assert final.status_code == 200
        document = Document(io.BytesIO(final.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert "포스코 강화 요구사항만 적용한다." in text
        assert "연결 KCS" not in text
        assert "내부 검토 메모" not in text
        assert "원본:" not in text

        review = client.get(f"/api/projects/{project_id}/export/review")
        review_document = Document(io.BytesIO(review.content))
        review_text = "\n".join(paragraph.text for paragraph in review_document.paragraphs)
        assert "간소화 초안" not in review_text
        assert "연결 KCS" not in review_text
        assert "내부 검토 메모" not in review_text
        assert "포스코 강화 요구사항만 적용한다." in review_text


def test_legacy_decision_without_reason_cannot_pass_final_export(tmp_path, monkeypatch):
    store, project_id = _store_with_project(tmp_path)
    store.update_decision(
        project_id,
        "clause",
        "keep",
        "포스코 기준",
        "",
        "safe-candidate",
        decision_reason="posco_specific",
    )
    with store.connect() as connection:
        connection.execute(
            "UPDATE clauses SET decision_reason = '' WHERE id = 'clause'"
        )
    project = store.get_project(project_id)
    assert project is not None
    assert project["invalid_decision_count"] == 1
    assert project["reviewed_clauses"] == 0
    assert project["final_export_ready"] is False

    monkeypatch.setattr(app_module, "store", store)
    monkeypatch.setattr(
        app_module,
        "_reconcile_current_kcs_revision",
        lambda: ("2026-09-04T22:34:56+09:00", {"revision": "test-revision"}),
    )
    with TestClient(app_module.app) as client:
        response = client.get(f"/api/projects/{project_id}/export/final")
    assert response.status_code == 409
    assert "판정 사유 미완료 1개" in response.json()["detail"]


def test_partial_overlap_keep_requires_a_nonempty_reduced_clause(tmp_path):
    store, project_id = _store_with_project(tmp_path)

    with pytest.raises(ValueError, match="문구를 입력"):
        store.update_decision(
            project_id,
            "clause",
            "keep",
            "",
            "",
            "safe-candidate",
            decision_reason="partial_overlap_residual",
        )
    with pytest.raises(ValueError, match="중복 내용을 제거"):
        store.update_decision(
            project_id,
            "clause",
            "keep",
            "강구조 요구사항을 적용한다.",
            "",
            "safe-candidate",
            decision_reason="partial_overlap_residual",
        )

    saved = store.update_decision(
        project_id,
        "clause",
        "keep",
        "포스코 추가 검사만 적용한다.",
        "",
        "safe-candidate",
        decision_reason="partial_overlap_residual",
    )

    assert saved is not None
    assert saved["edited_content"] == "포스코 추가 검사만 적용한다."


def test_final_export_is_blocked_when_later_ai_analysis_invalidates_safe_delete(tmp_path, monkeypatch):
    store, project_id = _store_with_project(tmp_path)
    store.save_coverage_analysis(
        project_id,
        "clause",
        ["safe-candidate", "warning-candidate"],
        {
            "coverage_status": "fully_covered",
            "confidence": 0.96,
            "requirements": [
                {
                    "requirement": "강구조 요구사항을 적용한다.",
                    "source_segment_ids": ["S1"],
                    "status": "covered",
                    "evidence_candidate_ids": ["safe-candidate"],
                    "evidence": "KCS에 같은 의무가 있다.",
                }
            ],
            "residual_content": "",
            "rationale": "전체 요구사항이 포함됩니다.",
        },
        "gpt-test",
    )
    store.update_decision(
        project_id,
        "clause",
        "delete",
        "",
        "",
        "safe-candidate",
        decision_reason="fully_covered_by_kcs",
        coverage_confirmed=True,
    )
    assert store.get_project(project_id)["review_submission_ready"] is True

    store.save_candidate_analysis(
        project_id,
        "clause",
        "safe-candidate",
        {
            "relation_type": "conflict",
            "confidence": 0.99,
            "rationale": "적용 조건이 충돌합니다.",
            "simplified_content": "포스코 적용 조건을 유지한다.",
        },
        "gpt-test",
    )
    monkeypatch.setattr(app_module, "store", store)
    monkeypatch.setattr(
        app_module,
        "_reconcile_current_kcs_revision",
        lambda: ("2026-09-04T22:34:56+09:00", {"revision": "test-revision"}),
    )

    with TestClient(app_module.app) as client:
        response = client.get(f"/api/projects/{project_id}/export/final")

    assert response.status_code == 409
    assert "삭제 근거 미충족 1개" in response.json()["detail"]


def test_legacy_coverage_without_source_binding_is_invalidated_idempotently(tmp_path):
    store, project_id = _store_with_project(tmp_path)
    with store.connect() as connection:
        connection.execute(
            """
            UPDATE clauses
            SET decision = 'delete', decision_reason = 'fully_covered_by_kcs',
                coverage_confirmed = 1, selected_candidate_id = 'safe-candidate'
            WHERE id = 'clause'
            """
        )
        connection.execute(
            """
            INSERT INTO clause_coverage_analysis(
                clause_id, candidate_ids_json, evidence_candidate_ids_json,
                coverage_status, confidence, requirements_json,
                residual_content, rationale, deletion_safe, model, analyzed_at
            ) VALUES ('clause', ?, ?, 'fully_covered', 0.99, ?, '',
                      '기존 분석', 1, 'legacy-model', '2026-01-01')
            """,
            (
                json.dumps(["safe-candidate", "warning-candidate"]),
                json.dumps(["safe-candidate"]),
                json.dumps(
                    [
                        {
                            "requirement": "요약된 기존 요구사항",
                            "status": "covered",
                            "evidence_candidate_ids": ["safe-candidate"],
                            "evidence": "기존 근거",
                        }
                    ],
                    ensure_ascii=False,
                ),
            ),
        )
        connection.execute(
            """
            INSERT INTO decision_history(
                clause_id, previous_decision, new_decision, edited_content,
                review_note, decision_reason, coverage_confirmed,
                selected_candidate_id, coverage_analysis_json, changed_at
            ) VALUES ('clause', NULL, 'delete', '', '',
                      'fully_covered_by_kcs', 1, 'safe-candidate', ?, '2026-01-01')
            """,
            (
                json.dumps(
                    {
                        "coverage_status": "fully_covered",
                        "deletion_safe": True,
                        "rationale": "기존 분석",
                    },
                    ensure_ascii=False,
                ),
            ),
        )

    migrated = Store(store.database_path)
    clause = migrated.get_clause(project_id, "clause")
    assert clause is not None
    assert clause["coverage_confirmed"] is False
    assert clause["coverage_analysis"]["coverage_status"] == "uncertain"
    assert clause["coverage_analysis"]["deletion_safe"] is False
    assert clause["coverage_analysis"]["requirements"][0]["source_segment_ids"] == ["S1"]
    first_requirements = clause["coverage_analysis"]["requirements"]
    with migrated.connect() as connection:
        history = connection.execute(
            """
            SELECT coverage_confirmed, coverage_analysis_json
            FROM decision_history
            WHERE clause_id = 'clause'
            ORDER BY id
            """
        ).fetchall()
    assert [bool(row["coverage_confirmed"]) for row in history] == [True, False]
    assert json.loads(history[-1]["coverage_analysis_json"])["coverage_status"] == "uncertain"

    migrated_again = Store(store.database_path)
    clause_again = migrated_again.get_clause(project_id, "clause")
    assert clause_again["coverage_analysis"]["requirements"] == first_requirements
    project = migrated_again.get_project(project_id)
    assert project["unsafe_delete_count"] == 1
    assert project["final_export_ready"] is False
    with migrated_again.connect() as connection:
        history_count = connection.execute(
            "SELECT COUNT(*) FROM decision_history WHERE clause_id = 'clause'"
        ).fetchone()[0]
    assert history_count == 2
