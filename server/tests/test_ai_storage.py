from __future__ import annotations

import io
import json

from openpyxl import load_workbook
import pytest

from server.documents import build_audit_xlsx
from server.storage import Store


def _project_with_candidate(tmp_path):
    store = Store(tmp_path / "ai-storage.sqlite3")
    project_id = "project-ai"
    clause_id = "clause-ai"
    candidate_id = "candidate-ai"
    store.create_project(
        {
            "id": project_id,
            "title": "AI 분석 시험",
            "source_filename": "source.docx",
            "source_path": str(tmp_path / "source.docx"),
            "source_sha256": "abc123",
            "uploaded_at": "2026-09-05T00:00:00+00:00",
            "kcs_snapshot": "2026-09-05T00:00:00+09:00",
            "kcs_revision": "test-revision",
            "kcs_scope": "KCS 41 31",
        },
        [
            {
                "id": clause_id,
                "source_order": 1,
                "label": "1.1",
                "title": "표면처리",
                "content": "강재 표면의 이물질을 제거하여야 한다.",
                "source_type": "paragraph",
                "outline_level": None,
                "candidates": [
                    {
                        "id": candidate_id,
                        "rank": 1,
                        "kcs_code": "KCS 41 31 05",
                        "document_name": "강구조공사 일반사항",
                        "version": "2026",
                        "update_date": "2026-08-01",
                        "kcs_clause": "3.2.1",
                        "title": "표면처리",
                        "content": "강재 표면의 이물질은 제거한다.",
                        "score": 0.8,
                        "classification": "높은 일치",
                        "reasons": ["핵심어 일치"],
                        "warnings": [],
                    }
                ],
            }
        ],
    )
    return store, project_id, clause_id, candidate_id


def test_ai_analysis_is_attached_to_visible_candidate(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)

    clause = store.save_candidate_analysis(
        project_id,
        clause_id,
        candidate_id,
        {
            "relation_type": "posco_specific",
            "confidence": 0.91,
            "rationale": "KCS에 없는 회사 관리 절차가 포함되어 있습니다.",
            "simplified_content": "회사 관리 절차는 유지한다.",
        },
        "gpt-test",
    )

    assert clause is not None
    analysis = clause["candidates"][0]["ai_analysis"]
    assert analysis["relation_type"] == "posco_specific"
    assert analysis["confidence"] == 0.91
    assert analysis["simplified_content"] == "회사 관리 절차는 유지한다."


def test_unrelated_ai_candidate_is_removed_from_all_visible_counts(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)

    clause = store.save_candidate_analysis(
        project_id,
        clause_id,
        candidate_id,
        {
            "relation_type": "unrelated",
            "confidence": 0.98,
            "rationale": "시공 단계와 검사 단계로 주제가 다릅니다.",
            "simplified_content": "",
        },
        "gpt-test",
    )

    assert clause is not None and clause["candidates"] == []
    assert store.list_clauses(project_id)[0]["candidate_count"] == 0
    project = store.get_project(project_id)
    assert project is not None
    assert project["candidate_clauses"] == 0
    assert project["high_match_clauses"] == 0


def test_low_confidence_unrelated_stays_visible_for_human_review(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)

    clause = store.save_candidate_analysis(
        project_id,
        clause_id,
        candidate_id,
        {
            "relation_type": "unrelated",
            "confidence": 0.7,
            "rationale": "무관할 가능성이 있지만 확신이 낮습니다.",
            "simplified_content": "",
        },
        "gpt-test",
    )

    assert clause is not None
    assert [item["id"] for item in clause["candidates"]] == [candidate_id]
    assert clause["excluded_candidates"] == []


def test_human_selected_candidate_cannot_be_hidden_by_gpt(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    store.update_decision(
        project_id, clause_id, "hold", "회사 기준", "", candidate_id,
        decision_reason="needs_expert_review",
    )

    try:
        store.save_candidate_analysis(
            project_id,
            clause_id,
            candidate_id,
            {
                "relation_type": "unrelated",
                "confidence": 0.99,
                "rationale": "GPT 판단",
                "simplified_content": "",
            },
            "gpt-test",
        )
    except ValueError as exc:
        assert "이미 선택한 후보" in str(exc)
    else:
        raise AssertionError("human-selected candidate must not be hidden")

    clause = store.get_clause(project_id, clause_id)
    assert clause is not None
    assert clause["selected_candidate_id"] == candidate_id
    assert clause["candidates"][0]["ai_analysis"] is None


def test_quality_ground_truth_candidate_cannot_be_hidden_by_gpt(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    store.ensure_quality_evaluation(project_id)
    store.update_quality_item(
        project_id,
        clause_id,
        "candidate_selected",
        candidate_id,
        "",
        "",
        "",
    )

    try:
        store.save_candidate_analysis(
            project_id,
            clause_id,
            candidate_id,
            {
                "relation_type": "unrelated",
                "confidence": 0.99,
                "rationale": "GPT 판단",
                "simplified_content": "",
            },
            "gpt-test",
        )
    except ValueError as exc:
        assert "품질평가 표본" in str(exc)
    else:
        raise AssertionError("quality ground truth candidate must not be hidden")

    evaluation = store.get_quality_evaluation(project_id)
    assert evaluation is not None
    assert evaluation["metrics"]["recall_at_3"] == 1.0


def test_unreviewed_quality_sample_candidate_set_is_locked(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    store.ensure_quality_evaluation(project_id)

    try:
        store.save_candidate_analysis(
            project_id,
            clause_id,
            candidate_id,
            {
                "relation_type": "unrelated",
                "confidence": 0.99,
                "rationale": "GPT 판단",
                "simplified_content": "",
            },
            "gpt-test",
        )
    except ValueError as exc:
        assert "후보군은 고정" in str(exc)
    else:
        raise AssertionError("quality sample candidates must remain fixed")


def test_excluded_candidate_cannot_be_restored_after_quality_sampling(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    store.save_candidate_analysis(
        project_id,
        clause_id,
        candidate_id,
        {
            "relation_type": "unrelated",
            "confidence": 0.99,
            "rationale": "GPT 판단",
            "simplified_content": "",
        },
        "gpt-test",
    )
    store.ensure_quality_evaluation(project_id)

    try:
        store.clear_candidate_analysis(project_id, clause_id, candidate_id)
    except ValueError as exc:
        assert "후보군은 고정" in str(exc)
    else:
        raise AssertionError("quality sample candidate set must not be restored")


def test_ai_analysis_rejects_invalid_relation(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)

    try:
        store.save_candidate_analysis(
            project_id,
            clause_id,
            candidate_id,
            {"relation_type": "low", "confidence": 0.5, "rationale": "invalid"},
            "gpt-test",
        )
    except ValueError as exc:
        assert "지원하지 않는" in str(exc)
    else:
        raise AssertionError("invalid relation must be rejected")


def test_selected_candidate_ai_analysis_is_exported_to_audit(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    store.save_candidate_analysis(
        project_id,
        clause_id,
        candidate_id,
        {
            "relation_type": "partial_overlap",
            "confidence": 0.87,
            "rationale": "검사 방법 일부만 일치합니다.",
            "simplified_content": "회사 추가 검사만 유지한다.",
        },
        "gpt-test",
    )
    store.update_decision(
        project_id, clause_id, "keep", "회사 추가 검사만 유지한다.", "", candidate_id,
        decision_reason="partial_overlap_residual",
    )

    project = store.get_project(project_id)
    assert project is not None
    workbook = load_workbook(io.BytesIO(build_audit_xlsx(project, store.audit_rows(project_id))))
    sheet = workbook["판정이력"]
    headers = [cell.value for cell in sheet[2]]
    values = {headers[index]: sheet.cell(3, index + 1).value for index in range(len(headers))}

    assert values["GPT 관계"] == "부분 일치"
    assert values["GPT 신뢰도"] == 0.87
    assert values["GPT 판정근거"] == "검사 방법 일부만 일치합니다."
    assert values["GPT 간소화 제안"] == "회사 추가 검사만 유지한다."


def test_audit_does_not_attach_current_candidate_to_history_without_one(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    store.update_decision(
        project_id,
        clause_id,
        "keep",
        "포스코 고유 기준만 유지한다.",
        "",
        None,
        decision_reason="posco_specific",
    )
    with store.connect() as connection:
        connection.execute(
            "UPDATE clauses SET selected_candidate_id = ? WHERE id = ?",
            (candidate_id, clause_id),
        )

    row = store.audit_rows(project_id)[0]

    assert row["kcs_code"] is None
    assert row["ai_relation_type"] is None


def test_audit_freezes_selected_candidate_ai_analysis_at_decision_time(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    store.save_candidate_analysis(
        project_id,
        clause_id,
        candidate_id,
        {
            "relation_type": "partial_overlap",
            "confidence": 0.87,
            "rationale": "일부만 일치합니다.",
            "simplified_content": "포스코 추가 기준만 유지한다.",
        },
        "gpt-before",
    )
    store.update_decision(
        project_id,
        clause_id,
        "keep",
        "포스코 추가 기준만 유지한다.",
        "",
        candidate_id,
        decision_reason="partial_overlap_residual",
    )
    store.save_candidate_analysis(
        project_id,
        clause_id,
        candidate_id,
        {
            "relation_type": "conflict",
            "confidence": 0.99,
            "rationale": "나중 분석은 충돌입니다.",
            "simplified_content": "",
        },
        "gpt-after",
    )

    row = store.audit_rows(project_id)[0]

    assert row["ai_relation_type"] == "partial_overlap"
    assert row["ai_confidence"] == 0.87
    assert row["ai_rationale"] == "일부만 일치합니다."


def test_partial_combined_coverage_preserves_residual_and_blocks_delete(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    clause = store.save_coverage_analysis(
        project_id,
        clause_id,
        [candidate_id],
        {
            "coverage_status": "partially_covered",
            "confidence": 0.96,
            "requirements": [
                {
                    "requirement": "강재 표면의 이물질을 제거하여야 한다.",
                    "source_segment_ids": ["S1"],
                    "status": "not_covered",
                    "evidence_candidate_ids": [],
                    "evidence": "KCS가 포스코 원문 전체를 포괄하지 않는다.",
                },
            ],
            "residual_content": "강재 표면의 염분 농도를 측정하여야 한다.",
            "rationale": "일부 요구사항만 KCS가 포함합니다.",
        },
        "gpt-test",
    )

    assert clause is not None
    assert clause["coverage_analysis"]["deletion_safe"] is False
    assert clause["coverage_analysis"]["residual_content"] == "강재 표면의 염분 농도를 측정하여야 한다."
    try:
        store.update_decision(
            project_id,
            clause_id,
            "delete",
            "",
            "",
            candidate_id,
            decision_reason="fully_covered_by_kcs",
            coverage_confirmed=True,
        )
    except ValueError as exc:
        assert "전체포괄" in str(exc)
    else:
        raise AssertionError("partial combined coverage must block deletion")


def test_candidate_reanalysis_invalidates_stale_combined_coverage(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    store.save_coverage_analysis(
        project_id,
        clause_id,
        [candidate_id],
        {
            "coverage_status": "fully_covered",
            "confidence": 0.96,
            "requirements": [
                {
                    "requirement": "표면 이물질을 제거한다.",
                    "source_segment_ids": ["S1"],
                    "status": "covered",
                    "evidence_candidate_ids": [candidate_id],
                    "evidence": "KCS에 같은 의무가 있다.",
                }
            ],
            "residual_content": "",
            "rationale": "전체 요구사항이 포함됩니다.",
        },
        "gpt-test",
    )
    store.save_candidate_analysis(
        project_id,
        clause_id,
        candidate_id,
        {
            "relation_type": "partial_overlap",
            "confidence": 0.9,
            "rationale": "개별 후보를 다시 분석했습니다.",
            "simplified_content": "추가 조건을 유지한다.",
        },
        "gpt-test",
    )

    clause = store.get_clause(project_id, clause_id)
    assert clause is not None
    assert clause["coverage_analysis"]["coverage_status"] == "uncertain"
    assert clause["coverage_analysis"]["deletion_safe"] is False
    assert "다시 실행" in clause["coverage_analysis"]["rationale"]


def test_restoring_excluded_candidate_cannot_erase_unsafe_coverage_gate(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    store.save_coverage_analysis(
        project_id,
        clause_id,
        [candidate_id],
        {
            "coverage_status": "partially_covered",
            "confidence": 0.97,
            "requirements": [
                {
                    "requirement": "포스코 추가 검사를 수행한다.",
                    "source_segment_ids": ["S1"],
                    "status": "not_covered",
                    "evidence_candidate_ids": [],
                    "evidence": "KCS에 추가 검사가 없다.",
                }
            ],
            "residual_content": "포스코 추가 검사를 수행하여야 한다.",
            "rationale": "포스코 고유 요구사항이 남습니다.",
        },
        "gpt-test",
    )
    store.save_candidate_analysis(
        project_id,
        clause_id,
        candidate_id,
        {
            "relation_type": "unrelated",
            "confidence": 0.99,
            "rationale": "서로 다른 검사입니다.",
            "simplified_content": "",
        },
        "gpt-test",
    )
    restored = store.clear_candidate_analysis(project_id, clause_id, candidate_id)

    assert restored is not None
    assert restored["coverage_analysis"]["coverage_status"] == "uncertain"
    assert restored["coverage_analysis"]["deletion_safe"] is False
    try:
        store.update_decision(
            project_id,
            clause_id,
            "delete",
            "",
            "",
            candidate_id,
            decision_reason="fully_covered_by_kcs",
            coverage_confirmed=True,
        )
    except ValueError as exc:
        assert "전체포괄" in str(exc)
    else:
        raise AssertionError("restoring a candidate must not remove an unsafe coverage gate")


def test_individual_conflict_blocks_later_fully_covered_analysis(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    store.save_candidate_analysis(
        project_id,
        clause_id,
        candidate_id,
        {
            "relation_type": "conflict",
            "confidence": 0.99,
            "rationale": "의무 조건이 서로 충돌합니다.",
            "simplified_content": "충돌 조건을 검토한다.",
        },
        "gpt-test",
    )
    clause = store.save_coverage_analysis(
        project_id,
        clause_id,
        [candidate_id],
        {
            "coverage_status": "fully_covered",
            "confidence": 0.99,
            "requirements": [
                {
                    "requirement": "표면 이물질을 제거한다.",
                    "source_segment_ids": ["S1"],
                    "status": "covered",
                    "evidence_candidate_ids": [candidate_id],
                    "evidence": "KCS에 같은 의무가 있다.",
                }
            ],
            "residual_content": "",
            "rationale": "문장만 보면 모두 포함됩니다.",
        },
        "gpt-test",
    )

    assert clause is not None
    assert clause["coverage_analysis"]["deletion_safe"] is False
    try:
        store.update_decision(
            project_id,
            clause_id,
            "delete",
            "",
            "",
            candidate_id,
            decision_reason="fully_covered_by_kcs",
            coverage_confirmed=True,
        )
    except ValueError as exc:
        assert "개별분석" in str(exc)
    else:
        raise AssertionError("an individual conflict must override later combined coverage")


def test_safe_delete_is_invalidated_by_later_conflicting_candidate_analysis(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    store.save_coverage_analysis(
        project_id,
        clause_id,
        [candidate_id],
        {
            "coverage_status": "fully_covered",
            "confidence": 0.96,
            "requirements": [
                {
                    "requirement": "표면 이물질을 제거한다.",
                    "source_segment_ids": ["S1"],
                    "status": "covered",
                    "evidence_candidate_ids": [candidate_id],
                    "evidence": "KCS에 같은 의무가 있다.",
                }
            ],
            "residual_content": "",
            "rationale": "전체 요구사항이 포함됩니다.",
        },
        "gpt-test",
    )
    store.update_decision(
        project_id,
        clause_id,
        "delete",
        "",
        "",
        candidate_id,
        decision_reason="fully_covered_by_kcs",
        coverage_confirmed=True,
    )
    assert store.get_project(project_id)["review_submission_ready"] is True

    clause = store.save_candidate_analysis(
        project_id,
        clause_id,
        candidate_id,
        {
            "relation_type": "conflict",
            "confidence": 0.99,
            "rationale": "수치 조건이 충돌합니다.",
            "simplified_content": "충돌 수치를 유지한다.",
        },
        "gpt-test",
    )
    project = store.get_project(project_id)

    assert clause is not None and clause["coverage_confirmed"] is False
    assert clause["coverage_analysis"]["coverage_status"] == "uncertain"
    assert project["unsafe_delete_count"] == 1
    assert project["final_export_ready"] is False
    audit_rows = store.audit_rows(project_id)
    confirmation_resets = [
        row
        for row in audit_rows
        if row["previous_decision"] == "delete"
        and row["new_decision"] == "delete"
        and row["coverage_confirmed"] == 0
    ]
    assert len(confirmation_resets) == 1
    assert confirmation_resets[0]["coverage_status"] == "uncertain"
    assert "무효화" in confirmation_resets[0]["coverage_rationale"]


def test_safe_coverage_reanalysis_requires_fresh_human_confirmation(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    analysis = {
        "coverage_status": "fully_covered",
        "confidence": 0.96,
        "requirements": [
            {
                "requirement": "표면 이물질을 제거한다.",
                "source_segment_ids": ["S1"],
                "status": "covered",
                "evidence_candidate_ids": [candidate_id],
                "evidence": "KCS에 같은 의무가 있다.",
            }
        ],
        "residual_content": "",
        "rationale": "전체 요구사항이 포함됩니다.",
    }
    store.save_coverage_analysis(
        project_id, clause_id, [candidate_id], analysis, "gpt-before"
    )
    store.update_decision(
        project_id,
        clause_id,
        "delete",
        "",
        "",
        candidate_id,
        decision_reason="fully_covered_by_kcs",
        coverage_confirmed=True,
    )
    assert store.get_project(project_id)["review_submission_ready"] is True

    refreshed = store.save_coverage_analysis(
        project_id, clause_id, [candidate_id], analysis, "gpt-after"
    )
    project = store.get_project(project_id)

    assert refreshed is not None and refreshed["coverage_confirmed"] is False
    assert project is not None and project["final_export_ready"] is False
    assert "삭제 근거 미충족 1개" in project["final_export_blockers"]


def test_combined_coverage_is_frozen_in_decision_history_and_audit(tmp_path):
    store, project_id, clause_id, first_candidate_id = _project_with_candidate(tmp_path)
    second_candidate_id = "candidate-ai-2"
    with store.connect() as connection:
        connection.execute(
            "UPDATE clauses SET content = ? WHERE id = ?",
            (
                "강재 표면의 이물질을 제거하여야 한다.\n염분 농도를 측정하여야 한다.",
                clause_id,
            ),
        )
        connection.execute(
            """
            INSERT INTO candidates(
                id, clause_id, rank, kcs_code, document_name, version,
                update_date, kcs_clause, title, content, score,
                classification, reasons_json, warnings_json
            ) VALUES (?, ?, 2, 'KCS 41 31 10', '강구조 검사기준', '2026',
                      '2026-08-02', '3.4.2', '염분 검사', ?, 0.77,
                      '관련성 높음', '[]', '[]')
            """,
            (second_candidate_id, clause_id, "강재 표면의 염분 농도를 측정한다."),
        )

    safe_analysis = {
        "coverage_status": "fully_covered",
        "confidence": 0.96,
        "requirements": [
            {
                "requirement": "요약문은 저장 시 원문으로 교체된다.",
                "source_segment_ids": ["S1"],
                "status": "covered",
                "evidence_candidate_ids": [first_candidate_id],
                "evidence": "이물질 제거 의무가 있다.",
            },
            {
                "requirement": "이 요약문도 원문으로 교체된다.",
                "source_segment_ids": ["S2"],
                "status": "covered",
                "evidence_candidate_ids": [second_candidate_id],
                "evidence": "염분 측정 의무가 있다.",
            },
        ],
        "residual_content": "",
        "rationale": "두 KCS를 조합하면 포스코 원문 전체를 포괄합니다.",
    }
    saved = store.save_coverage_analysis(
        project_id,
        clause_id,
        [first_candidate_id, second_candidate_id],
        safe_analysis,
        "gpt-history-test",
    )
    assert saved is not None
    assert saved["coverage_analysis"]["requirements"][0]["requirement"].startswith(
        "강재 표면의 이물질"
    )
    store.update_decision(
        project_id,
        clause_id,
        "delete",
        "",
        "복수 KCS 전체포괄 확인",
        first_candidate_id,
        decision_reason="fully_covered_by_kcs",
        coverage_confirmed=True,
    )
    with store.connect() as connection:
        history_before = connection.execute(
            "SELECT coverage_analysis_json FROM decision_history WHERE clause_id = ?",
            (clause_id,),
        ).fetchone()[0]
    snapshot = json.loads(history_before)
    assert snapshot["snapshot_version"] == 1
    assert snapshot["coverage_status"] == "fully_covered"
    assert snapshot["confidence"] == 0.96
    assert snapshot["model"] == "gpt-history-test"
    assert {item["id"] for item in snapshot["evidence_candidates"]} == {
        first_candidate_id,
        second_candidate_id,
    }
    assert snapshot["evidence_candidates"][0]["content"]

    store.save_coverage_analysis(
        project_id,
        clause_id,
        [first_candidate_id, second_candidate_id],
        {
            "coverage_status": "partially_covered",
            "confidence": 0.7,
            "requirements": [
                {
                    "requirement": "변경 후 첫 구간",
                    "source_segment_ids": ["S1"],
                    "status": "covered",
                    "evidence_candidate_ids": [first_candidate_id],
                    "evidence": "첫 구간만 포함한다.",
                },
                {
                    "requirement": "변경 후 둘째 구간",
                    "source_segment_ids": ["S2"],
                    "status": "not_covered",
                    "evidence_candidate_ids": [],
                    "evidence": "둘째 구간은 없다.",
                },
            ],
            "residual_content": "염분 농도를 측정하여야 한다.",
            "rationale": "재분석 결과 일부만 포함됩니다.",
        },
        "gpt-history-test-2",
    )
    with store.connect() as connection:
        history_after = connection.execute(
            "SELECT coverage_analysis_json FROM decision_history WHERE clause_id = ? ORDER BY id",
            (clause_id,),
        ).fetchall()
    assert len(history_after) == 2
    assert history_after[0][0] == history_before
    assert json.loads(history_after[1][0])["coverage_status"] == "partially_covered"

    project = store.get_project(project_id)
    current_clause = store.get_clause(project_id, clause_id)
    rows = store.audit_rows(project_id)
    assert project is not None and len(rows) == 2
    assert current_clause is not None and current_clause["coverage_confirmed"] is False
    assert project["final_export_ready"] is False
    assert rows[0]["coverage_status"] == "fully_covered"
    assert "KCS 41 31 05" in rows[0]["coverage_requirements"]
    assert "KCS 41 31 10" in rows[0]["coverage_requirements"]
    assert rows[1]["coverage_status"] == "partially_covered"
    workbook = load_workbook(io.BytesIO(build_audit_xlsx(project, rows)))
    headers = [cell.value for cell in workbook["판정이력"][2]]
    assert headers[:21] == [
        "순서", "포스코 조항", "포스코 제목", "포스코 원문", "현재 판정",
        "판정사유", "전체포괄 확인", "수정문", "검토의견", "검토일시", "선택 KCS",
        "KCS 문서명", "KCS 조항", "유사도", "GPT 관계", "GPT 신뢰도", "GPT 판정근거",
        "GPT 간소화 제안", "이전 판정", "변경 판정", "변경일시",
    ]
    assert "전체포괄 근거 KCS" in headers[21:]
    assert "요구사항별 분석" in headers[21:]


@pytest.mark.parametrize(
    "source_ids",
    (["S1"], ["S1", "S1"]),
    ids=("missing-segment", "duplicate-segment"),
)
def test_storage_rejects_incomplete_or_duplicate_source_segments(tmp_path, source_ids):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    with store.connect() as connection:
        connection.execute(
            "UPDATE clauses SET content = ? WHERE id = ?",
            ("첫 요구를 적용한다.\n둘째 요구를 적용한다.", clause_id),
        )
    requirements = [
        {
            "requirement": f"응답 {index}",
            "source_segment_ids": [source_id],
            "status": "covered",
            "evidence_candidate_ids": [candidate_id],
            "evidence": "KCS에 있다고 응답했다.",
        }
        for index, source_id in enumerate(source_ids, start=1)
    ]

    with pytest.raises(ValueError, match="누락하거나 중복"):
        store.save_coverage_analysis(
            project_id,
            clause_id,
            [candidate_id],
            {
                "coverage_status": "fully_covered",
                "confidence": 0.99,
                "requirements": requirements,
                "residual_content": "",
                "rationale": "전체라고 잘못 응답했다.",
            },
            "gpt-test",
        )


def _fully_covered_snapshot(candidate_id: str) -> dict:
    return {
        "coverage_status": "fully_covered",
        "confidence": 0.95,
        "requirements": [
            {
                "requirement": "강재 표면의 이물질을 제거하여야 한다.",
                "source_segment_ids": ["S1"],
                "status": "covered",
                "evidence_candidate_ids": [candidate_id],
                "evidence": "KCS에 같은 의무가 있다.",
            }
        ],
        "residual_content": "",
        "rationale": "전체 요구사항이 포함됩니다.",
    }


def test_coverage_save_rejects_source_changed_during_gpt_analysis(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    with store.connect() as connection:
        connection.execute(
            "UPDATE clauses SET content = ? WHERE id = ?",
            ("강재 표면의 염분을 측정하여야 한다.", clause_id),
        )

    with pytest.raises(ValueError, match="원문이 GPT 분석 중 변경"):
        store.save_coverage_analysis(
            project_id,
            clause_id,
            [candidate_id],
            _fully_covered_snapshot(candidate_id),
            "gpt-test",
            expected_posco_text="강재 표면의 이물질을 제거하여야 한다.",
        )

    with store.connect() as connection:
        saved_count = connection.execute(
            "SELECT COUNT(*) FROM clause_coverage_analysis WHERE clause_id = ?",
            (clause_id,),
        ).fetchone()[0]
    assert saved_count == 0


def test_coverage_save_rejects_candidate_content_changed_during_gpt_analysis(tmp_path):
    store, project_id, clause_id, candidate_id = _project_with_candidate(tmp_path)
    original_candidate_text = (
        "KCS 41 31 05\n강구조공사 일반사항\n3.2.1\n표면처리\n"
        "강재 표면의 이물질은 제거한다."
    )
    with store.connect() as connection:
        connection.execute(
            "UPDATE candidates SET content = ? WHERE id = ?",
            ("강재 표면의 염분은 측정한다.", candidate_id),
        )

    with pytest.raises(ValueError, match="KCS 후보 내용이 GPT 분석 중 변경"):
        store.save_coverage_analysis(
            project_id,
            clause_id,
            [candidate_id],
            _fully_covered_snapshot(candidate_id),
            "gpt-test",
            expected_posco_text="강재 표면의 이물질을 제거하여야 한다.",
            expected_candidate_texts=[original_candidate_text],
        )

    with store.connect() as connection:
        saved_count = connection.execute(
            "SELECT COUNT(*) FROM clause_coverage_analysis WHERE clause_id = ?",
            (clause_id,),
        ).fetchone()[0]
    assert saved_count == 0


@pytest.mark.parametrize("raw_snapshot", ("", "[]", "null", "{broken"))
def test_audit_tolerates_legacy_or_malformed_coverage_snapshot(tmp_path, raw_snapshot):
    store, project_id, clause_id, _candidate_id = _project_with_candidate(tmp_path)
    store.update_decision(
        project_id,
        clause_id,
        "keep",
        "",
        "",
        None,
        decision_reason="posco_specific",
    )
    with store.connect() as connection:
        connection.execute(
            "UPDATE decision_history SET coverage_analysis_json = ? WHERE clause_id = ?",
            (raw_snapshot, clause_id),
        )

    rows = store.audit_rows(project_id)

    assert len(rows) == 1
    assert rows[0]["coverage_status"] is None
    assert rows[0]["coverage_requirements"] == ""
