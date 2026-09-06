from __future__ import annotations

import hashlib
import io
import json
from datetime import datetime, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import server.app as app_module
from server.storage import (
    ClauseStructureConflictError,
    ReviewWorkflowConflictError,
    Store,
)


def _candidate() -> dict:
    return {
        "id": "workflow-candidate",
        "rank": 1,
        "kcs_code": "KCS 41 31 05",
        "document_name": "건축물 강구조공사 일반사항",
        "version": "2026",
        "update_date": "2026-09-01",
        "kcs_clause": "3.2.1",
        "title": "강구조 요구사항",
        "content": "강구조 요구사항을 적용한다.",
        "score": 0.82,
        "classification": "관련성 높음",
        "reasons": ["공통 핵심어"],
        "warnings": [],
    }


def _store_with_project(tmp_path, project_id: str = "workflow-project") -> tuple[Store, str]:
    store = Store(tmp_path / f"{project_id}.sqlite3")
    store.create_project(
        {
            "id": project_id,
            "title": "승인 워크플로 시험",
            "source_filename": "source.docx",
            "source_path": str(tmp_path / "source.docx"),
            "source_sha256": "source-hash",
            "uploaded_at": datetime.now(timezone.utc).isoformat(),
            "kcs_snapshot": "2026-09-04T22:34:56+09:00",
            "kcs_revision": "workflow-revision",
            "kcs_scope": "KCS 41 31",
            "warning": "",
        },
        [
            {
                "id": f"{project_id}-heading",
                "source_order": 1,
                "label": "1",
                "title": "일반사항",
                "content": "일반사항",
                "source_type": "heading",
                "outline_level": 1,
                "candidates": [],
            },
            {
                "id": f"{project_id}-clause",
                "source_order": 2,
                "label": "1.1",
                "title": "강구조 요구사항",
                "content": "강구조 요구사항을 적용한다.",
                "source_type": "paragraph",
                "outline_level": None,
                "candidates": [_candidate()],
            },
        ],
    )
    return store, project_id


def _mark_ready(store: Store, project_id: str, edited_content: str = "포스코 강화기준을 적용한다."):
    clause_id = f"{project_id}-clause"
    saved = store.update_decision(
        project_id,
        clause_id,
        "keep",
        edited_content,
        "",
        "workflow-candidate",
        decision_reason="posco_stricter",
    )
    assert saved is not None
    return saved


def _submission_row(store: Store, submission_id: str) -> dict:
    with store.connect() as connection:
        row = connection.execute(
            "SELECT * FROM project_review_submissions WHERE id = ?",
            (submission_id,),
        ).fetchone()
    assert row is not None
    return dict(row)


def test_submit_requires_ready_snapshot_then_locks_all_review_mutations(tmp_path):
    store, project_id = _store_with_project(tmp_path)

    initial = store.get_project(project_id)
    assert initial is not None
    assert initial["review_submission_ready"] is False
    assert initial["final_export_ready"] is False
    with pytest.raises(ReviewWorkflowConflictError, match="미검토 조항 1개"):
        store.submit_review(project_id, "작성자", "1차 검토 완료")

    _mark_ready(store, project_id)
    ready = store.get_project(project_id)
    assert ready is not None
    assert ready["review_submission_ready"] is True
    assert ready["final_export_ready"] is False

    workflow = store.submit_review(project_id, " 작성자 ", " 1차 검토 완료 ")
    submission = workflow["latest_review_submission"]
    assert workflow["status"] == "submitted"
    assert submission["revision_no"] == 1
    assert submission["author_name"] == "작성자"
    assert submission["author_note"] == "1차 검토 완료"

    stored = _submission_row(store, submission["id"])
    assert hashlib.sha256(stored["review_snapshot_json"].encode("utf-8")).hexdigest() == stored[
        "review_snapshot_sha256"
    ]
    snapshot = json.loads(stored["review_snapshot_json"])
    assert snapshot["project"]["source_sha256"] == "source-hash"
    assert snapshot["clauses"][1]["decision"] == "keep"
    assert snapshot["clauses"][1]["edited_content"] == "포스코 강화기준을 적용한다."

    locked = store.get_project(project_id)
    assert locked is not None
    assert locked["review_locked"] is True
    assert locked["final_export_ready"] is False
    with pytest.raises(ReviewWorkflowConflictError, match="변경할 수 없습니다"):
        _mark_ready(store, project_id, "제출 후 변경")
    with pytest.raises(ReviewWorkflowConflictError, match="변경할 수 없습니다"):
        store.quick_update_decision(
            project_id,
            f"{project_id}-clause",
            {"decision": "hold", "decision_reason": "needs_expert_review"},
        )
    with pytest.raises(ReviewWorkflowConflictError, match="변경할 수 없습니다"):
        store.clear_candidate_analysis(
            project_id,
            f"{project_id}-clause",
            "workflow-candidate",
        )
    with pytest.raises(ClauseStructureConflictError, match="구조를 변경할 수 없습니다"):
        store.prepare_clause_split(
            project_id,
            f"{project_id}-clause",
            ["강구조 요구사항을 ", "적용한다."],
        )


def test_api_requires_distinct_approver_reason_then_allows_reedit_and_resubmit(
    tmp_path,
    monkeypatch,
):
    store, project_id = _store_with_project(tmp_path)
    _mark_ready(store, project_id)
    monkeypatch.setattr(app_module, "store", store)

    with TestClient(app_module.app) as client:
        submitted = client.post(
            f"/api/projects/{project_id}/review-workflow/submit",
            json={"author_name": "작성자", "note": "초안 완료"},
        )
        assert submitted.status_code == 200
        first_id = submitted.json()["workflow"]["latest_review_submission"]["id"]

        same_person = client.post(
            f"/api/projects/{project_id}/review-workflow/{first_id}/approve",
            json={"approver_name": " 작성자 ", "note": "자가 승인"},
        )
        assert same_person.status_code == 409
        assert "다른 승인자" in same_person.json()["detail"]

        missing_reason = client.post(
            f"/api/projects/{project_id}/review-workflow/{first_id}/request-changes",
            json={"approver_name": "승인자", "reason": "   "},
        )
        assert missing_reason.status_code == 400
        assert "반려 사유" in missing_reason.json()["detail"]

        rejected = client.post(
            f"/api/projects/{project_id}/review-workflow/{first_id}/request-changes",
            json={"approver_name": "승인자", "reason": "회사 기준 문구를 명확히 할 것"},
        )
        assert rejected.status_code == 200
        assert rejected.json()["workflow"]["status"] == "changes_requested"
        assert rejected.json()["project"]["review_locked"] is False

        edited = client.patch(
            f"/api/projects/{project_id}/clauses/{project_id}-clause",
            json={
                "decision": "keep",
                "decision_reason": "posco_stricter",
                "coverage_confirmed": False,
                "edited_content": "포스코 강화기준과 추가 검사를 적용한다.",
                "review_note": "반려사항 반영",
                "selected_candidate_id": "workflow-candidate",
            },
        )
        assert edited.status_code == 200

        resubmitted = client.post(
            f"/api/projects/{project_id}/review-workflow/submit",
            json={"author_name": "작성자", "note": "반려사항 반영 완료"},
        )
        assert resubmitted.status_code == 200
        second = resubmitted.json()["workflow"]["latest_review_submission"]
        assert second["revision_no"] == 2

        approved = client.post(
            f"/api/projects/{project_id}/review-workflow/{second['id']}/approve",
            json={"approver_name": "승인자", "note": "최종 승인"},
        )
        assert approved.status_code == 200
        assert approved.json()["workflow"]["status"] == "approved"
        assert approved.json()["project"]["final_export_ready"] is True

        history = client.get(f"/api/projects/{project_id}/review-workflow")
        assert history.status_code == 200
        workflow = history.json()["workflow"]
        assert [row["revision_no"] for row in workflow["submissions"]] == [2, 1]
        assert [row["event_type"] for row in workflow["events"]] == [
            "submitted",
            "changes_requested",
            "submitted",
            "approved",
        ]
        assert [row["actor_name"] for row in workflow["events"]] == [
            "작성자",
            "승인자",
            "작성자",
            "승인자",
        ]

        audit = client.get(f"/api/projects/{project_id}/audit.xlsx")
        assert audit.status_code == 200

    assert store.workflow_audit_rows(project_id) == workflow["events"]
    workbook = load_workbook(io.BytesIO(audit.content), read_only=True)
    assert workbook.sheetnames == ["판정이력", "KCS 개정영향", "승인이력"]
    approval_rows = list(workbook["승인이력"].iter_rows(min_row=3, values_only=True))
    assert [row[2] for row in approval_rows] == ["검토 제출", "반려", "검토 제출", "승인"]


def test_approval_and_final_bundle_reject_snapshot_hash_mismatch(tmp_path):
    store, project_id = _store_with_project(tmp_path)
    _mark_ready(store, project_id)
    workflow = store.submit_review(project_id, "작성자", "검토 완료")
    submission_id = workflow["latest_review_submission"]["id"]

    with store.connect() as connection:
        connection.execute(
            "UPDATE clauses SET review_note = '제출 뒤 몰래 변경' WHERE id = ?",
            (f"{project_id}-clause",),
        )
    with pytest.raises(ReviewWorkflowConflictError, match="제출 후 검토 내용이 변경"):
        store.approve_review(project_id, submission_id, "승인자", "승인")

    with store.connect() as connection:
        connection.execute(
            "UPDATE clauses SET review_note = '' WHERE id = ?",
            (f"{project_id}-clause",),
        )
    approved = store.approve_review(project_id, submission_id, "승인자", "승인")
    assert approved["status"] == "approved"

    bundle = store.final_export_bundle(project_id)
    assert bundle is not None
    project, clauses = bundle
    assert project["final_export_ready"] is True
    assert [row["edited_content"] for row in clauses] == ["포스코 강화기준을 적용한다."]

    with store.connect() as connection:
        connection.execute(
            "UPDATE clauses SET edited_content = '승인 뒤 몰래 변경' WHERE id = ?",
            (f"{project_id}-clause",),
        )
    tampered = store.final_export_bundle(project_id)
    assert tampered is not None
    assert tampered[0]["final_export_ready"] is False
    assert "승인 후 검토 내용 변경 감지" in tampered[0]["final_export_blockers"]


def test_kcs_revision_change_supersedes_locked_submission_and_records_history(tmp_path):
    store, project_id = _store_with_project(tmp_path)
    _mark_ready(store, project_id)
    workflow = store.submit_review(project_id, "작성자", "검토 완료")
    submission_id = workflow["latest_review_submission"]["id"]

    store.publish_kcs_revision("2026-09-06T00:00:00+09:00", "new-revision")

    changed = store.get_review_workflow(project_id)
    assert changed is not None
    assert changed["status"] == "changes_requested"
    assert changed["latest_review_submission"]["id"] == submission_id
    assert changed["latest_review_submission"]["status"] == "superseded"
    assert [row["event_type"] for row in changed["events"]] == [
        "submitted",
        "superseded",
    ]
    assert changed["events"][-1]["actor_name"] == "시스템"
    project = store.get_project(project_id)
    assert project is not None
    assert project["review_locked"] is False
    assert project["final_export_ready"] is False
