from __future__ import annotations

import io
import json
from datetime import datetime, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import server.app as app_module
from server.storage import Store


def _candidate(candidate_id: str, clause_id: str, rank: int, score: float) -> dict:
    return {
        "id": candidate_id,
        "clause_id": clause_id,
        "rank": rank,
        "kcs_code": "KCS 41 31 05",
        "document_name": "건축물 강구조공사 일반사항",
        "version": "2026",
        "update_date": "2026-08-01T12:00:00",
        "kcs_clause": f"3.2.{rank}",
        "title": f"후보 {rank}",
        "content": f"KCS 후보 본문 {rank}",
        "score": score,
        "classification": "검토 필요",
        "reasons": ["시험용 후보"],
        "warnings": [],
    }


def _clause(
    clause_id: str,
    source_order: int,
    source_type: str = "paragraph",
    candidates: list[dict] | None = None,
) -> dict:
    return {
        "id": clause_id,
        "source_order": source_order,
        "label": str(source_order),
        "title": f"시험 조항 {source_order}",
        "content": f"포스코 시방서 본문 {source_order}",
        "source_type": source_type,
        "outline_level": 1 if source_type == "heading" else None,
        "candidates": candidates or [],
    }


def _insert_candidate(connection, candidate: dict) -> None:
    connection.execute(
        """
        INSERT INTO candidates(
            id, clause_id, rank, kcs_code, document_name, version,
            update_date, kcs_clause, title, content, score,
            classification, reasons_json, warnings_json
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            candidate["id"],
            candidate["clause_id"],
            candidate["rank"],
            candidate["kcs_code"],
            candidate["document_name"],
            candidate["version"],
            candidate["update_date"],
            candidate["kcs_clause"],
            candidate["title"],
            candidate["content"],
            candidate["score"],
            candidate["classification"],
            json.dumps(candidate["reasons"], ensure_ascii=False),
            json.dumps(candidate["warnings"], ensure_ascii=False),
        ),
    )


@pytest.fixture
def quality_client(tmp_path, monkeypatch):
    test_store = Store(tmp_path / "quality.sqlite3")
    project_id = "quality-project"
    clauses = [
        _clause("heading", 1, "heading"),
        _clause(
            "high-clause",
            2,
            candidates=[
                _candidate("high-rank-1", "high-clause", 1, 0.81),
                _candidate("high-rank-2", "high-clause", 2, 0.70),
            ],
        ),
        _clause(
            "review-clause",
            3,
            candidates=[_candidate("review-rank-1", "review-clause", 1, 0.33)],
        ),
        _clause("no-candidate-clause", 4),
    ]
    test_store.create_project(
        {
            "id": project_id,
            "title": "품질평가 시험",
            "source_filename": "quality.docx",
            "source_path": str(tmp_path / "quality.docx"),
            "source_sha256": "quality-source-sha256",
            "uploaded_at": datetime.now(timezone.utc).isoformat(),
            "kcs_snapshot": "2026-09-04T10:00:00+09:00",
            "kcs_revision": "test-revision",
            "kcs_scope": "KCS 41 31",
            "warning": "",
        },
        clauses,
    )
    monkeypatch.setattr(app_module, "store", test_store)
    with TestClient(app_module.app) as client:
        yield client, test_store, project_id


def test_quality_evaluation_is_idempotent_and_supports_small_projects(quality_client):
    client, test_store, project_id = quality_client

    response = client.post(
        f"/api/projects/{project_id}/quality-evaluation",
        json={"reviewer_name": " 품질 담당자 "},
    )
    assert response.status_code == 200
    evaluation = response.json()["evaluation"]
    assert evaluation["reviewer_name"] == "품질 담당자"
    assert evaluation["target_sample_size"] == 50
    assert evaluation["population_size"] == 3
    assert evaluation["actual_sample_size"] == 3
    assert [item["cohort"] for item in evaluation["items"]] == [
        "high",
        "review",
        "no_candidate",
    ]
    assert evaluation["metrics"]["sample_size"] == 3
    assert evaluation["metrics"]["remaining_count"] == 3
    assert evaluation["metrics"]["complete"] is False

    same = client.get(f"/api/projects/{project_id}/quality-evaluation")
    assert same.status_code == 200
    assert same.json()["evaluation"]["id"] == evaluation["id"]
    with test_store.connect() as connection:
        count = connection.execute("SELECT COUNT(*) FROM quality_evaluations").fetchone()[0]
    assert count == 1


def test_quality_verdict_validation_metrics_and_business_decision_separation(quality_client):
    client, test_store, project_id = quality_client
    client.get(f"/api/projects/{project_id}/quality-evaluation")

    business_clause = test_store.update_decision(
        project_id,
        "high-clause",
        "hold",
        "사내 특화 기준",
        "업무 판정 메모",
        "high-rank-1",
        decision_reason="needs_expert_review",
    )
    assert business_clause and business_clause["decision"] == "hold"

    missing_candidate = client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/high-clause",
        json={"verdict": "candidate_selected"},
    )
    assert missing_candidate.status_code == 400

    wrong_clause_candidate = client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/high-clause",
        json={
            "verdict": "candidate_selected",
            "relevant_candidate_id": "review-rank-1",
        },
    )
    assert wrong_clause_candidate.status_code == 400

    selected = client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/high-clause",
        json={
            "verdict": "candidate_selected",
            "relevant_candidate_id": "high-rank-2",
            "expected_kcs_code": "삭제되어야 함",
        },
    )
    assert selected.status_code == 200
    assert selected.json()["clause"]["decision"] == "hold"
    selected_item = next(
        item
        for item in selected.json()["evaluation"]["items"]
        if item["clause_id"] == "high-clause"
    )
    assert selected_item["relevant_rank"] == 2
    assert selected_item["expected_kcs_code"] == ""

    candidate_cannot_be_no_match = client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/review-clause",
        json={"verdict": "no_candidate_correct"},
    )
    assert candidate_cannot_be_no_match.status_code == 400

    no_candidate_cannot_reject_candidates = client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/no-candidate-clause",
        json={"verdict": "all_candidates_incorrect"},
    )
    assert no_candidate_cannot_reject_candidates.status_code == 400

    correct_no_candidate = client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/no-candidate-clause",
        json={"verdict": "no_candidate_correct", "quality_note": "독자 기준"},
    )
    assert correct_no_candidate.status_code == 200

    missing_code = client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/review-clause",
        json={"verdict": "kcs_missing", "expected_kcs_code": "  "},
    )
    assert missing_code.status_code == 400

    completed = client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/review-clause",
        json={
            "verdict": "kcs_missing",
            "expected_kcs_code": " KCS 41 31 20 ",
            "expected_kcs_clause": " 3.4.1 ",
            "quality_note": "후보군에서 누락",
        },
    )
    assert completed.status_code == 200
    evaluation = completed.json()["evaluation"]
    assert evaluation["completed_at"] is not None
    assert evaluation["metrics"]["complete"] is True
    assert evaluation["metrics"]["recall_at_3"] == 0.5
    assert evaluation["metrics"]["mrr"] == 0.25
    assert evaluation["metrics"]["accuracy"] == 0.6667
    assert evaluation["insights"]["error_signals"]["kcs_missing_with_candidates"] == 1
    assert evaluation["insights"]["error_signals"]["kcs_missing_without_candidates"] == 0

    unchanged_business = test_store.get_clause(project_id, "high-clause")
    assert unchanged_business and unchanged_business["decision"] == "hold"
    assert unchanged_business["review_note"] == "업무 판정 메모"

    cleared = client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/review-clause",
        json={"verdict": None},
    )
    assert cleared.status_code == 200
    evaluation = cleared.json()["evaluation"]
    assert evaluation["completed_at"] is None
    assert evaluation["metrics"]["complete"] is False
    cleared_item = next(
        item for item in evaluation["items"] if item["clause_id"] == "review-clause"
    )
    assert cleared_item["expected_kcs_code"] == ""
    assert cleared_item["quality_note"] == ""


def test_quality_endpoints_return_clear_not_found_errors(quality_client):
    client, _, project_id = quality_client

    missing_project = client.get("/api/projects/does-not-exist/quality-evaluation")
    assert missing_project.status_code == 404

    missing_sample_item = client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/heading",
        json={"verdict": None},
    )
    assert missing_sample_item.status_code == 404


def test_quality_insights_collect_without_premature_conclusions(quality_client):
    client, _, project_id = quality_client

    response = client.get(f"/api/projects/{project_id}/quality-evaluation")

    assert response.status_code == 200
    insights = response.json()["evaluation"]["insights"]
    assert insights["status"] == "collecting"
    assert "평가를 시작" in insights["message"]
    assert insights["evaluated_count"] == 0
    assert all(group["evaluated_count"] == 0 for group in insights["cohorts"].values())
    assert all(group["sample_accuracy"] is None for group in insights["score_bands"].values())
    assert {key: group["sample_count"] for key, group in insights["cohorts"].items()} == {
        "high": 1,
        "review": 1,
        "no_candidate": 1,
    }
    assert {
        key: group["sample_count"] for key, group in insights["score_bands"].items()
    } == {
        "gte_0_45": 1,
        "0_35_to_0_45": 0,
        "0_25_to_0_35": 1,
        "no_candidate": 1,
    }
    assert insights["error_signals"]["total"] == 0
    assert insights["error_signals"]["rate"] is None
    assert "recommended_threshold" not in insights


def test_quality_insights_group_evaluated_items_and_count_error_signals(quality_client):
    client, _, project_id = quality_client
    client.get(f"/api/projects/{project_id}/quality-evaluation")
    client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/high-clause",
        json={
            "verdict": "candidate_selected",
            "relevant_candidate_id": "high-rank-2",
        },
    )
    client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/review-clause",
        json={"verdict": "all_candidates_incorrect"},
    )
    response = client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/no-candidate-clause",
        json={"verdict": "kcs_missing", "expected_kcs_code": "KCS 41 31 20"},
    )

    assert response.status_code == 200
    insights = response.json()["evaluation"]["insights"]
    assert insights["status"] == "complete"
    assert insights["evaluated_count"] == 3
    assert insights["cohorts"]["high"] == {
        "sample_count": 1,
        "evaluated_count": 1,
        "correct_count": 1,
        "sample_accuracy": 1.0,
        "verdict_counts": {
            "candidate_selected": 1,
            "all_candidates_incorrect": 0,
            "no_candidate_correct": 0,
            "kcs_missing": 0,
        },
    }
    assert insights["cohorts"]["review"]["verdict_counts"]["all_candidates_incorrect"] == 1
    assert insights["cohorts"]["no_candidate"]["verdict_counts"]["kcs_missing"] == 1
    assert insights["score_bands"]["gte_0_45"]["evaluated_count"] == 1
    assert insights["score_bands"]["0_35_to_0_45"]["evaluated_count"] == 0
    assert insights["score_bands"]["0_25_to_0_35"]["evaluated_count"] == 1
    assert insights["score_bands"]["no_candidate"]["evaluated_count"] == 1
    assert insights["error_signals"] == {
        "candidate_selected_rank_2": 1,
        "candidate_selected_rank_3": 0,
        "candidate_selected_rank_2_or_3": 1,
        "all_candidates_incorrect": 1,
        "kcs_missing": 1,
        "kcs_missing_with_candidates": 0,
        "kcs_missing_without_candidates": 1,
        "total": 3,
        "rate": 1.0,
    }


def test_quality_xlsx_downloads_before_any_verdict_and_is_idempotent(quality_client):
    client, test_store, project_id = quality_client

    first = client.get(f"/api/projects/{project_id}/quality-evaluation.xlsx")
    second = client.get(f"/api/projects/{project_id}/quality-evaluation.xlsx")

    assert first.status_code == 200
    assert second.status_code == 200
    assert first.headers["cache-control"] == "no-store"
    assert "KCS%EB%A7%A4%EC%B9%AD_%ED%92%88%EC%A7%88%ED%8F%89%EA%B0%80.xlsx" in first.headers[
        "content-disposition"
    ]
    workbook = load_workbook(io.BytesIO(first.content))
    assert workbook.sheetnames == ["요약", "표본평가"]
    summary = workbook["요약"]
    assert summary["A1"].value == "KCS 매칭 품질평가 요약"
    summary_values = {
        summary.cell(row, 1).value: summary.cell(row, 2).value
        for row in range(2, 16)
    }
    assert summary_values["평가 상태"] == "평가 전"
    assert summary_values["평가 완료"] == 0
    assert any(cell.value == "판정 수" for cell in summary["A"])
    assert any(cell.value == "코호트별 통계" for cell in summary["A"])
    assert any(cell.value == "점수 구간별 통계" for cell in summary["A"])
    assert any(cell.value == "오류 신호" for cell in summary["A"])
    assert any(cell.value == "표본 수" for row in summary.iter_rows() for cell in row)

    detail = workbook["표본평가"]
    assert detail.max_row == 4
    headers = [cell.value for cell in detail[1]]
    assert "포스코 원문" in headers
    assert "후보3 score" in headers
    assert "선택후보순위" in headers
    assert detail.cell(2, headers.index("후보1 KCS 코드") + 1).value == "KCS 41 31 05"
    assert detail.cell(2, headers.index("평가결과") + 1).value is None

    with test_store.connect() as connection:
        evaluation_count = connection.execute(
            "SELECT COUNT(*) FROM quality_evaluations WHERE project_id = ?", (project_id,)
        ).fetchone()[0]
        item_count = connection.execute(
            """
            SELECT COUNT(*) FROM quality_evaluation_items qi
            JOIN quality_evaluations qe ON qe.id = qi.evaluation_id
            WHERE qe.project_id = ?
            """,
            (project_id,),
        ).fetchone()[0]
    assert evaluation_count == 1
    assert item_count == 3


def test_quality_sample_candidates_remain_locked(quality_client):
    client, test_store, project_id = quality_client
    client.get(f"/api/projects/{project_id}/quality-evaluation")

    with pytest.raises(ValueError, match="표본의 후보군은 고정"):
        test_store.save_candidate_analysis(
            project_id,
            "high-clause",
            "high-rank-1",
            {
                "relation_type": "unrelated",
                "confidence": 0.9,
                "rationale": "시험용 무관 판정",
                "simplified_content": "",
            },
            "test-model",
        )

    evaluation = test_store.ensure_quality_evaluation(project_id)
    assert evaluation and evaluation["actual_sample_size"] == 3
    high = next(item for item in evaluation["items"] if item["clause_id"] == "high-clause")
    assert high["candidate_count"] == 2

    restore = client.delete(
        f"/api/projects/{project_id}/clauses/high-clause/candidates/high-rank-1/ai-analysis"
    )
    assert restore.status_code == 400
    assert "표본의 후보군은 고정" in restore.json()["detail"]


def test_quality_metrics_and_export_keep_frozen_candidates_after_live_deletion(
    quality_client,
):
    client, test_store, project_id = quality_client
    started = client.get(f"/api/projects/{project_id}/quality-evaluation")
    assert started.status_code == 200

    with test_store.connect() as connection:
        connection.execute("DELETE FROM candidates WHERE clause_id = ?", ("high-clause",))

    selected = client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/high-clause",
        json={
            "verdict": "candidate_selected",
            "relevant_candidate_id": "high-rank-2",
        },
    )

    assert selected.status_code == 200
    evaluation = selected.json()["evaluation"]
    item = next(
        item for item in evaluation["items"] if item["clause_id"] == "high-clause"
    )
    assert item["candidate_count"] == 2
    assert item["top_score"] == pytest.approx(0.81)
    assert item["relevant_candidate_id"] is None
    assert item["relevant_rank"] == 2
    assert evaluation["metrics"]["rank_hits"]["2"] == 1
    assert evaluation["metrics"]["mrr"] == 0.5

    export_item = next(
        item
        for item in test_store.quality_export_rows(project_id)
        if item["clause_id"] == "high-clause"
    )
    assert [candidate["id"] for candidate in export_item["candidates"]] == [
        "high-rank-1",
        "high-rank-2",
    ]
    assert export_item["candidates"][0]["title"] == "후보 1"
    with test_store.connect() as connection:
        stored = connection.execute(
            """
            SELECT relevant_candidate_id, relevant_candidate_key,
                   relevant_rank_snapshot
            FROM quality_evaluation_items
            WHERE clause_id = ?
            """,
            ("high-clause",),
        ).fetchone()
    assert stored["relevant_candidate_id"] is None
    assert stored["relevant_candidate_key"]
    assert stored["relevant_rank_snapshot"] == 2


def test_empty_quality_snapshot_does_not_gain_later_live_candidates(quality_client):
    client, test_store, project_id = quality_client
    client.get(f"/api/projects/{project_id}/quality-evaluation")
    with test_store.connect() as connection:
        _insert_candidate(
            connection,
            _candidate("late-candidate", "no-candidate-clause", 1, 0.92),
        )

    test_store.initialize()
    response = client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/no-candidate-clause",
        json={"verdict": "no_candidate_correct"},
    )

    assert response.status_code == 200
    item = next(
        item
        for item in response.json()["evaluation"]["items"]
        if item["clause_id"] == "no-candidate-clause"
    )
    assert item["candidate_count"] == 0
    export_item = next(
        item
        for item in test_store.quality_export_rows(project_id)
        if item["clause_id"] == "no-candidate-clause"
    )
    assert export_item["candidates"] == []


def test_quality_schema_migration_backfills_candidates_and_selected_rank(quality_client):
    client, test_store, project_id = quality_client
    client.get(f"/api/projects/{project_id}/quality-evaluation")
    selected = client.patch(
        f"/api/projects/{project_id}/quality-evaluation/items/high-clause",
        json={
            "verdict": "candidate_selected",
            "relevant_candidate_id": "high-rank-2",
        },
    )
    assert selected.status_code == 200

    with test_store.connect() as connection:
        connection.execute("PRAGMA foreign_keys = OFF")
        connection.executescript(
            """
            CREATE TABLE quality_evaluation_items_legacy (
                evaluation_id TEXT NOT NULL REFERENCES quality_evaluations(id) ON DELETE CASCADE,
                clause_id TEXT NOT NULL REFERENCES clauses(id) ON DELETE CASCADE,
                sample_order INTEGER NOT NULL,
                cohort TEXT NOT NULL,
                verdict TEXT CHECK(verdict IN (
                    'candidate_selected', 'all_candidates_incorrect',
                    'no_candidate_correct', 'kcs_missing'
                ) OR verdict IS NULL),
                relevant_candidate_id TEXT REFERENCES candidates(id) ON DELETE SET NULL,
                expected_kcs_code TEXT NOT NULL DEFAULT '',
                expected_kcs_clause TEXT NOT NULL DEFAULT '',
                quality_note TEXT NOT NULL DEFAULT '',
                evaluated_at TEXT,
                updated_at TEXT,
                PRIMARY KEY(evaluation_id, clause_id),
                UNIQUE(evaluation_id, sample_order)
            );
            INSERT INTO quality_evaluation_items_legacy(
                evaluation_id, clause_id, sample_order, cohort, verdict,
                relevant_candidate_id, expected_kcs_code, expected_kcs_clause,
                quality_note, evaluated_at, updated_at
            )
            SELECT evaluation_id, clause_id, sample_order, cohort, verdict,
                   relevant_candidate_id, expected_kcs_code, expected_kcs_clause,
                   quality_note, evaluated_at, updated_at
            FROM quality_evaluation_items;
            DROP TABLE quality_evaluation_items;
            ALTER TABLE quality_evaluation_items_legacy
            RENAME TO quality_evaluation_items;
            """
        )

    migrated_store = Store(test_store.database_path)
    with migrated_store.connect() as connection:
        columns = {
            row["name"]
            for row in connection.execute("PRAGMA table_info(quality_evaluation_items)")
        }
        stored = connection.execute(
            """
            SELECT candidates_snapshot_json, relevant_candidate_key,
                   relevant_rank_snapshot
            FROM quality_evaluation_items
            WHERE clause_id = ?
            """,
            ("high-clause",),
        ).fetchone()
        connection.execute("DELETE FROM candidates WHERE id = ?", ("high-rank-2",))

    assert {
        "candidates_snapshot_json",
        "relevant_candidate_key",
        "relevant_rank_snapshot",
    } <= columns
    assert [candidate["id"] for candidate in json.loads(stored["candidates_snapshot_json"])] == [
        "high-rank-1",
        "high-rank-2",
    ]
    assert stored["relevant_candidate_key"]
    assert stored["relevant_rank_snapshot"] == 2
    evaluation = migrated_store.get_quality_evaluation(project_id)
    assert evaluation
    item = next(
        item for item in evaluation["items"] if item["clause_id"] == "high-clause"
    )
    assert item["relevant_candidate_id"] is None
    assert item["relevant_rank"] == 2
    assert evaluation["metrics"]["rank_hits"]["2"] == 1


def test_quality_xlsx_sanitizes_formula_control_characters_and_long_text(quality_client):
    client, test_store, project_id = quality_client
    dangerous_content = "=1+1\x01" + ("가" * 40000)
    with test_store.connect() as connection:
        connection.execute(
            "UPDATE projects SET title = ? WHERE id = ?",
            ("=HYPERLINK(\"https://example.test\")", project_id),
        )
        connection.execute(
            "UPDATE clauses SET content = ? WHERE id = ?",
            (dangerous_content, "high-clause"),
        )
        connection.execute(
            "UPDATE candidates SET title = ? WHERE id = ?",
            ("+SUM(1,1)\x02", "high-rank-1"),
        )

    response = client.get(f"/api/projects/{project_id}/quality-evaluation.xlsx")

    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content), data_only=False)
    assert workbook["요약"]["B2"].value.startswith("'=")
    detail = workbook["표본평가"]
    headers = [cell.value for cell in detail[1]]
    content = detail.cell(2, headers.index("포스코 원문") + 1).value
    candidate_title = detail.cell(2, headers.index("후보1 제목") + 1).value
    assert content.startswith("'=")
    assert "\x01" not in content
    assert len(content) == 32767
    assert candidate_title == "'+SUM(1,1)"
