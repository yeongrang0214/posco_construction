import io
import json
import uuid
from types import SimpleNamespace

from docx import Document
from docx.oxml import OxmlElement

from server.documents import parse_docx
from server import table_evidence as evidence


def source_row(grade="SM 275B", standard="KS D 3515"):
    return {"id": str(uuid.uuid4()), "source_type": "table", "content": f"1종 | B | {grade} | 강판, 형강 두께 200 ㎜ 이하",
            "match_context": f"2 재료 > (2) 용접구조용 압연 강재 ({standard}) > [표 열] 종별 | 재질 | 적용"}


def record(kind="KCS", code="413110"):
    return {"code": code, "codeType": kind, "name": "강구조 기준", "version": "2022", "updateDate": "2022-10-11",
            "list": [{"title": "구조용강재", "label": "표 2.1-1", "contents": '<table><tr><td>KS D 3515</td><td>SM275A, B, C, D, -TMC, SM460B, C</td></tr><tr><td>KS D 3503</td><td>SS275</td></tr></table>'}]}


def test_floating_table_rows_keep_caption_and_merged_conditions():
    doc = Document()
    doc.add_paragraph("2. 재료")
    doc.add_paragraph("2.1 강재")
    doc.add_paragraph("(1) 일반구조용 강재 (KS D 3503)")
    doc.add_paragraph("(2) 용접구조용 압연 강재 (KS D 3515)")
    anchor = doc.add_paragraph()
    table = doc.add_table(rows=3, cols=3)
    for i, values in enumerate([["종별", "재질", "적용"], ["A", "SM275A", "200 mm 이하"], ["B", "SM275B", ""]]):
        for cell, value in zip(table.rows[i].cells, values):
            cell.text = value
    table.cell(1, 2).merge(table.cell(2, 2))
    textbox = OxmlElement("w:txbxContent")
    textbox.append(table._tbl)
    anchor._p.append(textbox)
    doc.add_paragraph("(3) 다음 재료")
    other = doc.add_table(rows=1, cols=2)
    other.cell(0, 0).text = "SM275B"
    other.cell(0, 1).text = "SM275B"
    stream = io.BytesIO()
    doc.save(stream)
    _, rows, _ = parse_docx(stream.getvalue(), "test.docx", str(uuid.uuid4()))
    material = [r for r in rows if r["source_type"] == "table"]
    assert len(material) == 3
    assert all("200 mm 이하" in r["content"] for r in material[:2])
    assert all(evidence.source_fields(r)["standards"] == ["KS D 3515"] for r in material[:2])
    assert evidence.source_fields(material[-1]) is None
    assert material[-1]["content"] == "SM275B | SM275B"
    assert next(r for r in rows if r["content"] == "종별 | 재질 | 적용")["source_type"] == "heading"


def test_grade_suffix_expansion_never_substitutes_missing_grade():
    found = evidence.grades("SM275A, B, C, D, -TMC, SM355A, B, C, D, -TMC; SM460B, C")
    assert {"SM275B", "SM355D", "SM460C"} <= found
    assert "SM275-TMC" not in found
    assert "SM460A" not in found
    assert "SM275" not in found
    assert evidence.grades("SM 275A") == {"SM275A"}


def test_rowspan_reference_does_not_leak_into_next_ks_row():
    r = record()
    r["list"][0]["contents"] = '<table><tr><td rowspan="2">KS D 3515</td><td>SM275A</td></tr><tr><td>SM275B</td></tr><tr><td>KS D 3503</td><td>SS275</td></tr></table>'
    rows = evidence.record_rows(r, "KCS")
    assert rows[1]["standards"] == ["KS D 3515"]
    assert rows[2]["standards"] == ["KS D 3503"]


def test_no_parent_ks_inheritance_when_nearest_caption_has_no_ks():
    row = source_row()
    row["match_context"] = "KS D 3515 > 다른 강재"
    assert evidence.source_fields(row) is None


def test_actual_table_identifiers_do_not_prove_thickness(tmp_path, monkeypatch):
    raw = tmp_path / "raw"
    raw.mkdir()
    (raw / "KCS_413110.json").write_text(json.dumps([record()]), encoding="utf-8")
    settings = SimpleNamespace(kcs_raw_dir=raw)
    monkeypatch.setattr(evidence, "_design_record", lambda *_: {"record": record("KDS", "413010")})
    row = source_row()
    row["candidates"] = evidence.table_candidates(row, raw, ("4131",))
    result = evidence.compare_table(row, settings, ("4131",))
    assert [r["kind"] for r in result["evidence"]] == ["KCS", "KDS"]
    assert result["evidence"][0]["candidate_id"] == row["candidates"][0]["id"]
    assert result["evidence"][1]["candidate_id"] is None
    assert [c["status"] for c in result["checks"]] == ["found", "found", "unconfirmed", "unconfirmed"]
    assert row["candidates"][0]["classification"] == "일부 요구사항 대응"
    assert evidence.table_candidates(source_row("SM460A"), raw, ("4131",)) == []
    assert evidence.table_candidates(source_row(standard="KS D 3529"), raw, ("4131",)) == []
    # A second request must not reuse a different clause's candidate id.
    second = evidence.compare_table(source_row(), settings, ("4131",))
    assert second["evidence"][0]["candidate_id"] is None


def test_kds_failure_is_visible_and_does_not_break_kcs(tmp_path, monkeypatch):
    raw = tmp_path / "raw"
    raw.mkdir()
    (raw / "KCS_413110.json").write_text(json.dumps([record()]), encoding="utf-8")
    monkeypatch.setattr(evidence, "_design_record", lambda *_: {"warning": "KDS 조회 실패"})
    result = evidence.compare_table(source_row(), SimpleNamespace(kcs_raw_dir=raw), ("4131",))
    assert result["warnings"] == ["KDS 조회 실패"]
    assert len(result["evidence"]) == 1


def test_kds_catalog_revision_must_match_viewer(monkeypatch):
    monkeypatch.setattr(evidence, "_design_cache", {})
    monkeypatch.setattr(evidence, "_load_api_key", lambda _: "test-key")
    old = record("KDS", "413010")
    current = {**old, "version": "2099"}
    monkeypatch.setattr(evidence, "_request_json", lambda path, *_: [current] if path == "CodeList" else [old])
    result = evidence._design_record(None, "413010")
    assert "record" not in result
    assert "warning" in result


def test_kds_is_searched_even_without_kcs_match(tmp_path, monkeypatch):
    monkeypatch.setattr(evidence, "_design_record", lambda *_: {"record": record("KDS", "413010")})
    result = evidence.compare_table(source_row(), SimpleNamespace(kcs_raw_dir=tmp_path), ("4131",))
    assert [r["kind"] for r in result["evidence"]] == ["KDS"]
    assert result["evidence"][0]["candidate_id"] is None


def test_changed_kcs_body_is_not_linked_to_a_stale_saved_candidate(tmp_path, monkeypatch):
    (tmp_path / "KCS_413110.json").write_text(json.dumps([record()]), encoding="utf-8")
    monkeypatch.setattr(evidence, "_design_record", lambda *_: {"record": record("KDS", "413010")})
    row = source_row()
    row["candidates"] = evidence.table_candidates(row, tmp_path, ("4131",))
    row["candidates"][0]["content"] = "이전 개정의 표"
    result = evidence.compare_table(row, SimpleNamespace(kcs_raw_dir=tmp_path), ("4131",))
    assert result["evidence"][0]["candidate_id"] is None


def test_table_endpoint_decisions_and_export_do_not_require_notes(tmp_path, monkeypatch):
    from dataclasses import replace
    from fastapi.testclient import TestClient
    from server import app as application
    from server.storage import Store

    raw = tmp_path / "raw"
    raw.mkdir()
    (raw / "KCS_413110.json").write_text(json.dumps([record()]), encoding="utf-8")
    test_store = Store(tmp_path / "tables.sqlite3")
    project = str(uuid.uuid4())
    rows = []
    for index, grade in enumerate(("SM275A", "SM275B", "SM275C", "SM275D"), 1):
        row = {**source_row(grade), "title": "1종", "source_order": index, "label": f"표 1-{index}", "outline_level": None}
        row["candidates"] = evidence.table_candidates(row, raw, ("4131",))
        rows.append(row)
    test_store.create_project({"id": project, "title": "표 검토 시험", "source_filename": "건축_제13장_철골공사_201712.docx",
        "source_path": str(tmp_path / "source.docx"), "source_sha256": "test-hash", "uploaded_at": "2026-09-10T00:00:00Z",
        "kcs_snapshot": "2026-09-10", "kcs_revision": "table-test-revision", "kcs_scope": "KCS 41 31"}, rows)
    monkeypatch.setattr(application, "store", test_store)
    monkeypatch.setattr(application, "settings", replace(application.settings, kcs_raw_dir=raw, exports_dir=tmp_path))
    monkeypatch.setattr(evidence, "_design_record", lambda *_: {"record": record("KDS", "413010")})
    client = TestClient(application.app)
    base = f"/api/projects/{project}/clauses"
    response = client.get(f"{base}/{rows[0]['id']}/table-comparison")
    assert response.status_code == 200
    assert len(response.json()["evidence"]) == 2
    candidate_id = rows[0]["candidates"][0]["id"]
    assert client.patch(f"{base}/{rows[0]['id']}/quick-review", json={"selected_candidate_id": candidate_id}).status_code == 200
    response = client.patch(f"{base}/{rows[0]['id']}/quick-review", json={"decision": "keep", "decision_reason": "no_kcs_match", "selected_candidate_id": None})
    assert response.status_code == 200
    assert response.json()["clause"]["selected_candidate_id"] is None
    assert client.patch(f"{base}/{rows[1]['id']}/quick-review", json={"decision": "delete", "decision_reason": "management_decision"}).status_code == 200
    assert client.patch(f"{base}/{rows[2]['id']}/quick-review", json={"decision": "hold", "decision_reason": "needs_expert_review"}).status_code == 200
    exported = client.get(f"/api/projects/{project}/export/review")
    assert exported.status_code == 200, exported.text
    document = Document(io.BytesIO(exported.content))
    text = " ".join(p.text for p in document.paragraphs)
    assert all(grade in text for grade in ("SM275A", "SM275C", "SM275D"))
    assert "SM275B" not in text
    assert "KCS" not in text and "KDS" not in text
    assert "1종 1종" not in text
