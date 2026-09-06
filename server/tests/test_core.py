from __future__ import annotations

import io
import json
import sqlite3
import uuid
from datetime import datetime, timezone

from docx import Document
from openpyxl import load_workbook

from server.documents import build_audit_xlsx, build_review_docx, parse_docx, sha256_bytes
from server.matcher import match_clauses, read_snapshot
from server.storage import Store


def make_source_docx() -> bytes:
    document = Document()
    document.core_properties.title = "철골공사 시험 시방서"
    document.add_heading("1. 일반사항", level=1)
    document.add_paragraph("1.1 적용범위")
    document.add_paragraph("강재의 표면은 도장 전에 먼지, 유분 및 기타 이물질을 제거하여야 한다.")
    document.add_paragraph("자동번호 적용 문단", style="List Number")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "구분"
    table.cell(0, 1).text = "기준"
    table.cell(1, 0).text = "도장"
    table.cell(1, 1).text = "표면의 이물질을 제거한다."
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_kcs_snapshot(tmp_path):
    root = tmp_path / "kcs"
    raw = root / "raw"
    raw.mkdir(parents=True)
    manifest = {
        "status": "completed",
        "latestOnly": True,
        "completedAt": "2026-09-04T10:00:00+09:00",
        "documentCount": 1,
    }
    (root / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    code_list = [
        {
            "codeType": "KCS",
            "code": "413105",
            "fullCode": "KCS 41 31 05",
            "name": "건축물 강구조공사 일반사항",
            "version": "2026",
            "updateDate": "2026-08-01T12:00:00",
        }
    ]
    (root / "code_list.json").write_text(
        json.dumps(code_list, ensure_ascii=False), encoding="utf-8"
    )
    payload = [
        {
            "codeType": "KCS",
            "code": "413105",
            "fullCode": "KCS 41 31 05",
            "name": "건축물 강구조공사 일반사항",
            "version": "2026",
            "updateDate": "2026-08-01T12:00:00",
            "list": [
                {
                    "label": "3.2.1",
                    "title": "표면처리",
                    "contents": "<p>강재의 표면은 도장 전에 먼지, 유분 및 기타 이물질을 제거하여야 한다.</p>",
                }
            ],
        }
    ]
    (raw / "KCS_413105.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return root, raw


def test_parse_skips_toc_and_builds_contextual_labels():
    document = Document()
    document.add_paragraph("목 차")
    document.add_paragraph("1. 일반사항")
    document.add_paragraph("1.1 적용범위")
    document.add_paragraph("1. 일반사항")
    document.add_paragraph("1.2 제출물")
    document.add_paragraph("본문 첫 줄이다.")
    document.add_paragraph("(1) 작업계획")
    document.add_paragraph("(6) 시험계획")
    document.add_paragraph("가. 재료")
    output = io.BytesIO()
    document.save(output)

    project_id = str(uuid.uuid4())
    _, clauses, warnings = parse_docx(output.getvalue(), "context.docx", project_id)

    assert [clause["label"] for clause in clauses] == [
        "1", "1.2", "1.2(1)", "1.2(6)", "1.2(6)-가.",
    ]
    assert clauses[1]["content"] == "본문 첫 줄이다."
    assert clauses[1]["match_context"] == "1 일반사항"
    assert all("목 차" not in clause["title"] for clause in clauses)
    assert any("목차와 표제부" in warning for warning in warnings)


def test_parse_match_review_and_export(tmp_path):
    data = make_source_docx()
    project_id = str(uuid.uuid4())
    title, clauses, warnings = parse_docx(data, "건축_제13장_철골공사.docx", project_id)
    assert title == "철골공사 시험 시방서"
    assert len(clauses) == 5
    assert clauses[0]["source_type"] == "heading"
    assert clauses[1]["label"] == "1.1"
    assert "강재의 표면은" in clauses[1]["content"]
    assert any(clause["source_type"] == "table" for clause in clauses)
    assert warnings

    kcs_root, raw = make_kcs_snapshot(tmp_path)
    snapshot, manifest = read_snapshot(kcs_root / "manifest.json")
    assert snapshot.startswith("2026-09-04")
    assert manifest["latestOnly"] is True

    match_clauses(clauses, raw, ("4131",))
    matched = next(clause for clause in clauses if clause["candidates"])
    assert matched["candidates"][0]["kcs_code"] == "KCS 41 31 05"
    assert matched["candidates"][0]["score"] >= 0.45

    source_path = tmp_path / "source.docx"
    source_path.write_bytes(data)
    project = {
        "id": project_id,
        "title": title,
        "source_filename": source_path.name,
        "source_path": str(source_path),
        "source_sha256": sha256_bytes(data),
        "uploaded_at": datetime.now(timezone.utc).isoformat(),
        "kcs_snapshot": snapshot,
        "kcs_revision": manifest["revision"],
        "kcs_scope": "KCS 41 31",
        "warning": " ".join(warnings),
    }
    store = Store(tmp_path / "review.sqlite3")
    store.create_project(project, clauses)
    project_summary = store.get_project(project_id)
    assert project_summary and project_summary["candidate_clauses"] >= 1
    assert project_summary["high_match_clauses"] >= 1
    assert project_summary["table_clauses"] == 2
    selected = matched["candidates"][0]["id"]
    saved = store.update_decision(
        project_id, matched["id"], "hold", "포스코 추가 기준", "재검토", selected,
        decision_reason="needs_expert_review",
    )
    assert saved and saved["decision"] == "hold"
    assert store.get_project(project_id)["reviewed_clauses"] == 1

    export_rows = store.export_clauses(project_id, ("keep", "hold"))
    review_bytes = build_review_docx(project, export_rows, "review")
    review_doc = Document(io.BytesIO(review_bytes))
    assert "[보류]" in "\n".join(paragraph.text for paragraph in review_doc.paragraphs)
    assert "포스코 추가 기준" in "\n".join(paragraph.text for paragraph in review_doc.paragraphs)

    audit_bytes = build_audit_xlsx(project, store.audit_rows(project_id))
    workbook = load_workbook(io.BytesIO(audit_bytes))
    assert workbook["판정이력"].max_row >= 3


def test_legacy_database_migration_preserves_rows_and_is_idempotent(tmp_path):
    database_path = tmp_path / "legacy.sqlite3"
    with sqlite3.connect(database_path) as connection:
        connection.executescript(
            """
            CREATE TABLE projects (
                id TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                source_filename TEXT NOT NULL,
                source_path TEXT NOT NULL,
                source_sha256 TEXT NOT NULL,
                uploaded_at TEXT NOT NULL,
                kcs_snapshot TEXT NOT NULL,
                kcs_scope TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'reviewing',
                warning TEXT NOT NULL DEFAULT ''
            );
            CREATE TABLE clauses (
                id TEXT PRIMARY KEY,
                project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
                source_order INTEGER NOT NULL,
                label TEXT NOT NULL,
                title TEXT NOT NULL,
                content TEXT NOT NULL,
                source_type TEXT NOT NULL,
                outline_level INTEGER,
                decision TEXT,
                edited_content TEXT NOT NULL DEFAULT '',
                review_note TEXT NOT NULL DEFAULT '',
                selected_candidate_id TEXT,
                reviewed_at TEXT,
                UNIQUE(project_id, source_order)
            );
            CREATE TABLE decision_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                clause_id TEXT NOT NULL REFERENCES clauses(id) ON DELETE CASCADE,
                previous_decision TEXT,
                new_decision TEXT,
                edited_content TEXT NOT NULL,
                review_note TEXT NOT NULL,
                selected_candidate_id TEXT,
                changed_at TEXT NOT NULL
            );
            """
        )
        connection.execute(
            """
            INSERT INTO projects(
                id, title, source_filename, source_path, source_sha256,
                uploaded_at, kcs_snapshot, kcs_scope
            ) VALUES ('legacy-project', '구 시방서', 'legacy.docx', 'legacy.docx',
                      'legacy-sha', '2026-01-01', '2026-01-01', 'KCS 41')
            """
        )
        connection.execute(
            """
            INSERT INTO clauses(
                id, project_id, source_order, label, title, content, source_type,
                decision, edited_content, review_note, reviewed_at
            ) VALUES ('legacy-clause', 'legacy-project', 1, '1.1', '기존 조항',
                      '기존 내용', 'paragraph', 'keep', '기존 출력문', '기존 메모', '2026-01-02')
            """
        )
        connection.execute(
            """
            INSERT INTO decision_history(
                clause_id, previous_decision, new_decision, edited_content,
                review_note, changed_at
            ) VALUES ('legacy-clause', NULL, 'keep', '기존 출력문', '기존 메모', '2026-01-02')
            """
        )

    Store(database_path)
    Store(database_path)

    with sqlite3.connect(database_path) as connection:
        connection.row_factory = sqlite3.Row
        clause = connection.execute(
            "SELECT * FROM clauses WHERE id = 'legacy-clause'"
        ).fetchone()
        history = connection.execute(
            "SELECT * FROM decision_history WHERE clause_id = 'legacy-clause'"
        ).fetchone()
        legacy_project = connection.execute(
            "SELECT * FROM projects WHERE id = 'legacy-project'"
        ).fetchone()
        coverage_table = connection.execute(
            "SELECT name FROM sqlite_master WHERE type = 'table' AND name = 'clause_coverage_analysis'"
        ).fetchone()
        coverage_index = connection.execute(
            "SELECT name FROM sqlite_master WHERE type = 'index' AND name = 'idx_clause_coverage_status'"
        ).fetchone()

    assert clause is not None and clause["content"] == "기존 내용"
    assert clause["decision_reason"] == ""
    assert clause["coverage_confirmed"] == 0
    assert history is not None and history["review_note"] == "기존 메모"
    assert history["decision_reason"] == ""
    assert history["coverage_confirmed"] == 0
    assert history["coverage_analysis_json"] == ""
    assert legacy_project is not None and legacy_project["kcs_revision"] == ""
    assert coverage_table is not None
    assert coverage_index is not None
