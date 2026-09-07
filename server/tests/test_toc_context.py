from __future__ import annotations

import io
import uuid
import zipfile
from xml.etree import ElementTree

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

import server.matcher as matcher_module
from server.documents import parse_docx


def _docx_bytes(document: Document) -> bytes:
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_generic_word_core_title_falls_back_to_filename(tmp_path):
    document = Document()
    document.core_properties.title = "Word Document"
    document.add_paragraph("1. 일반사항")
    document.add_paragraph("이 시방서는 시험공사에 적용한다.")

    title, _clauses, _warnings = parse_docx(
        _docx_bytes(document), "전기_제01장_일반사항_210903.docx", str(uuid.uuid4())
    )

    assert title == "전기_제01장_일반사항_210903"


def test_parse_handles_document_without_numbering_part():
    document = Document()
    document.add_paragraph("1. 일반사항")
    document.add_paragraph("본 공사에 적용하는 기준을 따른다.")
    original = _docx_bytes(document)
    stripped = io.BytesIO()

    with zipfile.ZipFile(io.BytesIO(original), "r") as source:
        with zipfile.ZipFile(stripped, "w") as target:
            for item in source.infolist():
                if item.filename == "word/numbering.xml":
                    continue
                payload = source.read(item.filename)
                if item.filename == "word/_rels/document.xml.rels":
                    root = ElementTree.fromstring(payload)
                    for relationship in list(root):
                        if relationship.get("Type", "").endswith("/numbering"):
                            root.remove(relationship)
                    payload = ElementTree.tostring(
                        root, encoding="utf-8", xml_declaration=True
                    )
                target.writestr(item, payload)

    title, clauses, _warnings = parse_docx(
        stripped.getvalue(),
        "건축_제01장_일반사항_201712.docx",
        str(uuid.uuid4()),
    )

    assert title == "건축_제01장_일반사항_201712"
    assert clauses


def test_parse_skips_normal_style_toc_when_outline_restarts():
    document = Document()
    document.add_paragraph("철골공사 시방서")
    document.add_paragraph("1. 일반사항")
    document.add_paragraph("1.1 적용범위")
    document.add_paragraph("2. 재료")
    document.add_paragraph("2.1 강재")
    document.add_paragraph("3. 시공")
    document.add_paragraph("1. 일반사항")
    document.add_paragraph("1.1 적용범위")
    document.add_paragraph("이 시방서는 철골공사에 적용한다.")
    document.add_paragraph("2. 재료")
    document.add_paragraph("2.1 강재")
    document.add_paragraph("(1) 강재는 검사하여야 한다.")

    _, clauses, warnings = parse_docx(
        _docx_bytes(document), "normal-toc.docx", str(uuid.uuid4())
    )

    assert [clause["label"] for clause in clauses] == [
        "1",
        "1.1",
        "2",
        "2.1",
        "2.1(1)",
    ]
    assert clauses[-1]["match_context"] == "2 재료 > 2.1 강재"
    assert any("목차와 표제부" in warning for warning in warnings)


def test_parse_keeps_uncertain_short_repeated_outline():
    document = Document()
    document.add_paragraph("1. 일반사항")
    document.add_paragraph("1.1 적용범위")
    document.add_paragraph("본문 설명이다.")
    document.add_paragraph("1. 일반사항")
    document.add_paragraph("1.2 제출물")

    _, clauses, warnings = parse_docx(
        _docx_bytes(document), "not-a-toc.docx", str(uuid.uuid4())
    )

    assert [clause["label"] for clause in clauses] == ["1", "1.1", "1", "1.2"]
    assert not any("목차와 표제부" in warning for warning in warnings)


def test_parse_keeps_long_outline_when_only_first_anchor_repeats():
    document = Document()
    document.add_paragraph("1. 일반사항")
    document.add_paragraph("1.1 적용범위")
    document.add_paragraph("1.2 제출물")
    document.add_paragraph("2. 재료")
    document.add_paragraph("2.1 강재")
    document.add_paragraph("1. 일반사항")
    document.add_paragraph("1.1 적용범위")
    document.add_paragraph("1.3 시공계획")
    document.add_paragraph("작업계획을 제출하여야 한다.")

    _, clauses, warnings = parse_docx(
        _docx_bytes(document), "repeated-body.docx", str(uuid.uuid4())
    )

    assert [clause["label"] for clause in clauses] == [
        "1",
        "1.1",
        "1.2",
        "2",
        "2.1",
        "1",
        "1.1",
        "1.3",
    ]
    assert not any("목차와 표제부" in warning for warning in warnings)


def test_parse_continues_past_first_page_break_in_multi_page_toc():
    document = Document()
    document.add_paragraph("철골공사 시방서")
    try:
        toc_style = document.styles["TOC 1"]
    except KeyError:
        toc_style = document.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("1. 일반사항", style=toc_style)
    first_page_end = document.add_paragraph("1.1 적용범위", style=toc_style)
    first_page_end.add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("2. 재료", style=toc_style)
    document.add_paragraph("2.1 강재", style=toc_style)
    document.add_paragraph("1. 일반사항")
    document.add_paragraph("1.1 적용범위")
    document.add_paragraph("이 시방서는 철골공사에 적용한다.")

    _, clauses, warnings = parse_docx(
        _docx_bytes(document), "styled-toc.docx", str(uuid.uuid4())
    )

    assert [clause["label"] for clause in clauses] == ["1", "1.1"]
    assert clauses[1]["match_context"] == "1 일반사항"
    assert all(clause["title"] not in {"재료", "강재"} for clause in clauses)
    assert any("목차와 표제부" in warning for warning in warnings)


def test_parse_treats_plain_numeric_after_toc_page_break_as_body():
    document = Document()
    document.add_paragraph("철골공사 시방서")
    try:
        toc_style = document.styles["TOC 1"]
    except KeyError:
        toc_style = document.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("1. 일반사항", style=toc_style)
    last_toc_entry = document.add_paragraph("2. 재료", style=toc_style)
    last_toc_entry.add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("1. 총칙")
    document.add_paragraph("1.1 적용범위")
    document.add_paragraph("이 시방서는 철골공사에 적용한다.")

    _, clauses, warnings = parse_docx(
        _docx_bytes(document), "single-page-toc.docx", str(uuid.uuid4())
    )

    assert [clause["label"] for clause in clauses] == ["1", "1.1"]
    assert clauses[1]["match_context"] == "1 총칙"
    assert any("목차와 표제부" in warning for warning in warnings)


def test_bare_numeric_requirements_preserve_numbers_in_matcher_input(
    tmp_path,
    monkeypatch,
):
    document = Document()
    document.add_paragraph("90 m/m 이상으로 한다.")
    document.add_paragraph("100% 실시한다.")
    document.add_paragraph("2개 이상 설치한다.")
    document.add_paragraph("50 Kg 이상으로 한다.")
    document.add_paragraph("1회 이상 실시한다.")
    document.add_paragraph("3 개월 동안 보관한다.")

    _, clauses, _ = parse_docx(
        _docx_bytes(document), "numeric-requirements.docx", str(uuid.uuid4())
    )

    assert all(clause["label"] not in {"90", "100", "2", "50", "1", "3"} for clause in clauses)
    assert all(clause["source_type"] == "paragraph" for clause in clauses)
    preserved_text = " ".join(
        str(clause["content"] or clause["title"]) for clause in clauses
    )
    for expected in (
        "90 m/m",
        "100%",
        "2개",
        "50 Kg",
        "1회",
        "3 개월",
    ):
        assert expected in preserved_text

    matcher_inputs: list[str] = []

    def fake_rank_matches(source_text, _raw_dir, _prefixes):
        matcher_inputs.append(source_text)
        return []

    monkeypatch.setattr(matcher_module, "_rank_matches", fake_rank_matches)
    matcher_module.match_clauses(clauses, tmp_path, ("4131",))
    combined_matcher_input = " ".join(matcher_inputs)
    for expected in ("90 m/m", "100%", "2개", "50 Kg", "1회", "3 개월"):
        assert expected in combined_matcher_input


def test_short_numeric_noun_heading_stays_heading_without_continuation():
    document = Document()
    document.add_paragraph("19. 제품운반 및 현장야적")

    _, clauses, _ = parse_docx(
        _docx_bytes(document), "numeric-heading.docx", str(uuid.uuid4())
    )

    assert len(clauses) == 1
    assert clauses[0]["source_type"] == "heading"
    assert clauses[0]["title"] == "제품운반 및 현장야적"


def test_nested_short_titles_are_context_only_but_requirements_stay_reviewable():
    document = Document()
    document.add_paragraph("2.1 무석면 바닥타일 붙이기")
    document.add_paragraph("(1) 붙이기")
    document.add_paragraph("(2) 접착제는 바탕면 전체에 도포하여야 한다.")
    document.add_paragraph("(3) 응결시간은 10시간 이내")

    _, clauses, _ = parse_docx(
        _docx_bytes(document), "nested-headings.docx", str(uuid.uuid4())
    )

    assert [clause["source_type"] for clause in clauses] == [
        "heading",
        "heading",
        "paragraph",
        "paragraph",
    ]
    assert clauses[1]["match_context"].endswith("2.1 무석면 바닥타일 붙이기")


def test_bare_integer_heading_accepts_conservative_one_character_toc_difference():
    document = Document()
    try:
        toc_style = document.styles["TOC 1"]
    except KeyError:
        toc_style = document.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("14 포장", style=toc_style)
    last_toc_entry = document.add_paragraph("15 제품검사", style=toc_style)
    last_toc_entry.add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("15 완제품검사")

    _, clauses, warnings = parse_docx(
        _docx_bytes(document), "toc-backed-heading.docx", str(uuid.uuid4())
    )

    assert len(clauses) == 1
    assert clauses[0]["label"] == "15"
    assert clauses[0]["source_type"] == "heading"
    assert clauses[0]["title"] == "완제품검사"
    assert any("목차와 표제부" in warning for warning in warnings)


def test_heading_is_promoted_when_sentence_continuation_is_appended():
    document = Document()
    document.add_paragraph("19. 제품운반 및 현장야적")
    document.add_paragraph("제품은 손상되지 않도록 운반하여야 한다.")

    _, clauses, _ = parse_docx(
        _docx_bytes(document), "heading-continuation.docx", str(uuid.uuid4())
    )

    assert len(clauses) == 1
    assert clauses[0]["source_type"] == "paragraph"
    assert clauses[0]["content"] == "제품은 손상되지 않도록 운반하여야 한다."


def test_zero_leading_measurement_is_not_a_label_or_numeric_context():
    document = Document()
    document.add_paragraph("1.1 배합비")
    document.add_paragraph("0.108   2   0.216 L/M2")
    document.add_paragraph("마. 후속기준")

    _, clauses, _ = parse_docx(
        _docx_bytes(document), "decimal-measurement.docx", str(uuid.uuid4())
    )

    assert all(not clause["label"].startswith("0.108") for clause in clauses)
    assert "0.108 2 0.216 L/M2" in clauses[0]["content"]
    assert clauses[1]["label"] == "1.1-마."


def test_decimal_quantities_are_preserved_but_decimal_section_label_remains():
    document = Document()
    document.add_paragraph("1.5 m 이상으로 한다.")
    document.add_paragraph("12.5 mm 이하로 한다.")
    document.add_paragraph("0.8 MPa 이상이어야 한다.")
    document.add_paragraph("1.5 제출물")

    _, clauses, _ = parse_docx(
        _docx_bytes(document), "decimal-quantities.docx", str(uuid.uuid4())
    )

    assert clauses[-1]["label"] == "1.5"
    assert clauses[-1]["title"] == "제출물"
    quantity_text = " ".join(
        str(clause["content"] or clause["title"]) for clause in clauses[:-1]
    )
    assert "1.5 m" in quantity_text
    assert "12.5 mm" in quantity_text
    assert "0.8 MPa" in quantity_text
    assert all(
        clause["label"] not in {"12.5", "0.8"}
        for clause in clauses[:-1]
    )


def test_matcher_only_matches_paragraphs_and_tables_with_context(tmp_path, monkeypatch):
    ranked_source_texts: list[str] = []

    def fake_rank_matches(source_text, raw_dir, prefixes):
        ranked_source_texts.append(source_text)
        return []

    monkeypatch.setattr(matcher_module, "_rank_matches", fake_rank_matches)
    clauses = [
        {
            "id": str(uuid.uuid4()),
            "title": "일반사항",
            "content": "",
            "source_type": "heading",
        },
        {
            "id": str(uuid.uuid4()),
            "title": "목차 항목",
            "content": "",
            "source_type": "toc",
        },
        {
            "id": str(uuid.uuid4()),
            "title": "검사",
            "content": "강재를 검사한다.",
            "source_type": "paragraph",
            "match_context": "2 재료 > 2.1 강재",
        },
        {
            "id": str(uuid.uuid4()),
            "title": "구분",
            "content": "구분 | 기준",
            "source_type": "table",
            "match_context": "3 시공",
        },
    ]

    matcher_module.match_clauses(clauses, tmp_path, ("4131",))

    assert len(ranked_source_texts) == 2
    assert ranked_source_texts[0].startswith("2 재료 > 2.1 강재 검사 검사")
    assert ranked_source_texts[1].startswith("3 시공 구분 구분")
    assert clauses[0]["candidates"] == []
    assert clauses[1]["candidates"] == []
