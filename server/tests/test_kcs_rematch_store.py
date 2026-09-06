from __future__ import annotations

import json
from io import BytesIO
from copy import deepcopy
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from server.documents import build_audit_xlsx
from server.storage import (
    CURRENT_PARSER_VERSION,
    KCSRematchConflictError,
    Store,
)


def _candidate(candidate_id: str, clause_id: str, rank: int, *, score: float = 0.8) -> dict:
    return {
        "id": candidate_id,
        "clause_id": clause_id,
        "rank": rank,
        "kcs_code": "KCS 41 31 05",
        "document_name": "건축물 강구조공사 일반사항",
        "version": "2026",
        "update_date": "2026-08-01",
        "kcs_clause": f"3.2.{rank}",
        "title": f"후보 {rank}",
        "content": f"KCS 의미 본문 {rank}",
        "score": score,
        "classification": "관련성 높음",
        "reasons": ["시험 후보"],
        "warnings": [],
    }


def _clause(clause_id: str, order: int, candidates: list[dict] | None = None) -> dict:
    return {
        "id": clause_id,
        "source_order": order,
        "label": str(order),
        "title": f"조항 {order}",
        "content": f"포스코 본문 {order}",
        "source_type": "paragraph",
        "outline_level": None,
        "match_context": "철골공사 > 시험",
        "candidates": candidates or [],
    }


def _create_project(
    store: Store,
    tmp_path,
    *,
    project_id: str = "rematch-project",
    ready_decision: str | None = None,
) -> str:
    clauses = [
        {
            "id": "heading",
            "source_order": 1,
            "label": "1",
            "title": "철골공사",
            "content": "",
            "source_type": "heading",
            "outline_level": 1,
            "match_context": "철골공사",
            "candidates": [],
        },
        _clause(
            "reviewed",
            2,
            [
                _candidate("candidate-1", "reviewed", 1, score=0.82),
                _candidate("candidate-2", "reviewed", 2, score=0.71),
            ],
        ),
    ]
    store.create_project(
        {
            "id": project_id,
            "title": "재매칭 저장소 시험",
            "source_filename": "source.docx",
            "source_path": str(tmp_path / "source.docx"),
            "source_sha256": f"sha-{project_id}",
            "uploaded_at": datetime.now(timezone.utc).isoformat(),
            "kcs_snapshot": "snapshot-a",
            "kcs_revision": "revision-a",
            "kcs_scope": "KCS 41 31",
            "parser_version": CURRENT_PARSER_VERSION,
            "warning": "",
        },
        clauses,
    )
    if ready_decision:
        with store.connect() as connection:
            connection.execute(
                """
                UPDATE clauses
                SET decision = ?, decision_reason = 'posco_specific',
                    edited_content = '사내 잔존 문구', review_note = '기존 검토',
                    reviewed_at = '2026-09-01T00:00:00+00:00'
                WHERE id = 'reviewed'
                """,
                (ready_decision,),
            )
    return project_id


def _seed_delete_evidence(store: Store) -> None:
    with store.connect() as connection:
        connection.execute(
            """
            UPDATE clauses
            SET decision = 'delete', decision_reason = 'fully_covered_by_kcs',
                selected_candidate_id = 'candidate-1', coverage_confirmed = 1,
                edited_content = '기존 편집문', review_note = '기존 삭제 판정',
                reviewed_at = '2026-09-01T00:00:00+00:00'
            WHERE id = 'reviewed'
            """
        )
        connection.executemany(
            """
            INSERT INTO candidate_ai_analysis(
                candidate_id, relation_type, confidence, rationale,
                simplified_content, model, analyzed_at
            ) VALUES (?, ?, ?, ?, '', 'test-model', '2026-09-01')
            """,
            (
                ("candidate-1", "equivalent", 0.96, "동일 기준"),
                ("candidate-2", "partial_overlap", 0.75, "일부 기준"),
            ),
        )
        connection.execute(
            """
            INSERT INTO clause_coverage_analysis(
                clause_id, candidate_ids_json, evidence_candidate_ids_json,
                coverage_status, confidence, requirements_json,
                residual_content, rationale, deletion_safe, model, analyzed_at
            ) VALUES (
                'reviewed', '["candidate-1", "candidate-2"]', '["candidate-1"]',
                'fully_covered', 0.96, '[]', '', '전체 포괄', 1,
                'test-model', '2026-09-01'
            )
            """
        )


def _matcher_candidate(candidate: dict, **changes) -> dict:
    result = {
        key: deepcopy(candidate[key])
        for key in (
            "id",
            "rank",
            "kcs_code",
            "document_name",
            "version",
            "update_date",
            "kcs_clause",
            "title",
            "content",
            "score",
            "classification",
            "reasons",
            "warnings",
        )
    }
    result.update(changes)
    return result


def _results(prepared: dict) -> list[dict]:
    return [
        {
            "id": clause["id"],
            "candidates": [
                _matcher_candidate(candidate) for candidate in clause["candidates"]
            ],
        }
        for clause in prepared["clauses"]
    ]


def _queue_and_claim(store: Store, *, target: str = "revision-b") -> tuple[dict, dict]:
    runs = store.publish_kcs_revision(f"snapshot-{target[-1]}", target)
    assert len(runs) == 1
    run = runs[0]
    prepared = store.claim_kcs_rematch(run["id"], {"mode": "test"})
    return run, prepared


def test_publish_is_idempotent_and_same_revision_needs_no_run(tmp_path):
    store = Store(tmp_path / "queue.sqlite3")
    project_id = _create_project(store, tmp_path)

    assert store.publish_kcs_revision("snapshot-a2", "revision-a") == []
    assert store.list_kcs_rematch_runs(project_id) == []

    first = store.publish_kcs_revision("snapshot-b", "revision-b")
    second = store.publish_kcs_revision("snapshot-b", "revision-b")

    assert len(first) == len(second) == 1
    assert first[0]["id"] == second[0]["id"]
    assert first[0]["status"] == "pending"
    assert store.ensure_kcs_rematch(project_id)["id"] == first[0]["id"]
    assert len(store.list_kcs_rematch_runs(project_id)) == 1
    audit = store.kcs_rematch_audit_rows(project_id)
    assert len(audit) == 1
    assert audit[0]["run_id"] == first[0]["id"]
    assert audit[0]["clause_id"] is None


def test_legacy_parser_project_requires_reupload_and_is_not_queued(tmp_path):
    store = Store(tmp_path / "legacy-parser.sqlite3")
    project_id = _create_project(store, tmp_path)
    with store.connect() as connection:
        connection.execute(
            "UPDATE projects SET parser_version = '' WHERE id = ?",
            (project_id,),
        )

    assert store.publish_kcs_revision("snapshot-b", "revision-b") == []
    project = store.get_project(project_id)

    assert project["requires_source_reupload"] is True
    assert project["kcs_rematch"] is None
    assert "원본 시방서 재업로드 필요" in project["final_export_blockers"]
    with pytest.raises(KCSRematchConflictError, match="다시 업로드"):
        store.ensure_kcs_rematch(project_id)


def test_metadata_only_rematch_preserves_ids_ai_coverage_and_decision(tmp_path):
    store = Store(tmp_path / "metadata.sqlite3")
    project_id = _create_project(store, tmp_path)
    _seed_delete_evidence(store)
    run, prepared = _queue_and_claim(store)
    results = _results(prepared)
    reviewed = next(item for item in results if item["id"] == "reviewed")
    first, second = reviewed["candidates"]
    first.update(rank=2, score=0.55, version="2027", update_date="2027-01-01")
    second.update(rank=1, score=0.91, version="2027", update_date="2027-01-01")

    completed = store.apply_kcs_rematch(
        run["id"],
        prepared["expected_state_sha256"],
        results,
        {"mode": "test"},
    )

    assert completed["status"] == "completed"
    assert completed["material_change_count"] == 0
    assert completed["review_required_count"] == 0
    assert completed["matched_count"] == 1
    assert completed["impacts"][0]["impact_type"] == "metadata_change"
    assert completed["impacts"][0]["review_required"] is False
    clause = store.get_clause(project_id, "reviewed")
    assert [candidate["id"] for candidate in clause["candidates"]] == [
        "candidate-2",
        "candidate-1",
    ]
    assert clause["selected_candidate_id"] == "candidate-1"
    assert clause["decision"] == "delete"
    assert clause["coverage_confirmed"] is True
    assert clause["coverage_analysis"] is not None
    assert clause["coverage_analysis"]["candidate_ids"] == [
        "candidate-2",
        "candidate-1",
    ]
    with store.connect() as connection:
        ai_ids = {
            row[0]
            for row in connection.execute(
                "SELECT candidate_id FROM candidate_ai_analysis"
            ).fetchall()
        }
    assert ai_ids == {"candidate-1", "candidate-2"}
    assert store.get_project(project_id)["kcs_revision"] == "revision-b"


def test_material_change_preserves_decision_but_invalidates_evidence_and_selection(tmp_path):
    store = Store(tmp_path / "material.sqlite3")
    project_id = _create_project(store, tmp_path)
    _seed_delete_evidence(store)
    run, prepared = _queue_and_claim(store)
    results = _results(prepared)
    reviewed = next(item for item in results if item["id"] == "reviewed")
    reviewed["candidates"][0]["content"] = "개정되어 의미가 달라진 KCS 본문"

    completed = store.apply_kcs_rematch(
        run["id"], prepared["expected_state_sha256"], results
    )

    assert completed["material_change_count"] == 1
    assert completed["review_required_count"] == 1
    impact = completed["impacts"][0]
    assert impact["impact_type"] == "material_change"
    assert impact["review_required"] is True
    assert "candidate_meaning_changed" in impact["reasons"]
    assert impact["before_decision"]["decision"] == "delete"
    assert impact["before_decision"]["coverage_analysis"]["deletion_safe"] == 1
    clause = store.get_clause(project_id, "reviewed")
    assert clause["decision"] == "delete"
    assert clause["edited_content"] == "기존 편집문"
    assert clause["review_note"] == "기존 삭제 판정"
    assert clause["selected_candidate_id"] is None
    assert clause["coverage_confirmed"] is False
    assert clause["coverage_analysis"] is None
    assert clause["kcs_impact"]["run_id"] == run["id"]


def test_fingerprint_change_and_new_current_revision_leave_apply_atomic(tmp_path):
    store = Store(tmp_path / "atomic.sqlite3")
    project_id = _create_project(store, tmp_path)
    run, prepared = _queue_and_claim(store)
    results = _results(prepared)
    with store.connect() as connection:
        connection.execute(
            "UPDATE clauses SET review_note = '동시 변경' WHERE id = 'reviewed'"
        )

    with pytest.raises(KCSRematchConflictError, match="재매칭 중 변경"):
        store.apply_kcs_rematch(run["id"], prepared["expected_state_sha256"], results)

    assert store.get_project(project_id)["kcs_revision"] == "revision-a"
    assert [candidate["id"] for candidate in store.get_clause(project_id, "reviewed")["candidates"]] == [
        "candidate-1",
        "candidate-2",
    ]
    assert store.get_kcs_rematch_run(run["id"])["status"] == "running"

    store.publish_kcs_revision("snapshot-c", "revision-c")
    with pytest.raises(KCSRematchConflictError):
        store.apply_kcs_rematch(run["id"], prepared["expected_state_sha256"], results)
    assert store.get_kcs_rematch_run(run["id"])["status"] == "superseded"
    assert store.get_project(project_id)["kcs_revision"] == "revision-a"


def test_fail_and_retry_do_not_change_candidates_or_project_revision(tmp_path):
    store = Store(tmp_path / "retry.sqlite3")
    project_id = _create_project(store, tmp_path)
    run, prepared = _queue_and_claim(store)

    failed = store.fail_kcs_rematch(
        run["id"], "matcher failed", prepared["expected_state_sha256"]
    )
    retried = store.retry_kcs_rematch(run["id"])

    assert failed["status"] == "failed"
    assert failed["error"] == "matcher failed"
    assert retried["status"] == "pending"
    assert retried["expected_state_sha256"] == ""
    assert store.get_project(project_id)["kcs_revision"] == "revision-a"
    assert [candidate["id"] for candidate in store.get_clause(project_id, "reviewed")["candidates"]] == [
        "candidate-1",
        "candidate-2",
    ]
    reclaimed = store.claim_kcs_rematch(run["id"], {"mode": "test-2"})
    assert reclaimed["run"]["status"] == "running"


def test_ack_is_atomic_with_decision_and_blocks_stale_revision_race(tmp_path):
    store = Store(tmp_path / "ack.sqlite3")
    project_id = _create_project(store, tmp_path, ready_decision="keep")
    run, prepared = _queue_and_claim(store)
    results = _results(prepared)
    reviewed = next(item for item in results if item["id"] == "reviewed")
    reviewed["candidates"][0]["content"] = "새 KCS 의미"
    store.apply_kcs_rematch(run["id"], prepared["expected_state_sha256"], results)

    blocked = store.get_project(project_id)
    assert blocked["unacknowledged_kcs_impact_count"] == 1
    assert blocked["final_export_ready"] is False
    items, total = store.bulk_review(project_id, status="kcs_impact")
    assert total == 1
    assert items[0]["kcs_impact"]["run_id"] == run["id"]

    updated = store.update_decision(
        project_id,
        "reviewed",
        "keep",
        "사내 잔존 문구",
        "영향 재검토 완료",
        None,
        decision_reason="posco_specific",
        expected_kcs_revision="revision-b",
        impact_run_id=run["id"],
        acknowledge_kcs_impact=True,
    )
    assert updated["kcs_impact"]["acknowledged_at"] is not None
    assert store.get_project(project_id)["review_submission_ready"] is True

    store.publish_kcs_revision("snapshot-c", "revision-c")
    with pytest.raises(KCSRematchConflictError, match="개정이 변경"):
        store.update_decision(
            project_id,
            "reviewed",
            "keep",
            "사내 잔존 문구",
            "경합 중 저장되면 안 됨",
            None,
            decision_reason="posco_specific",
            expected_kcs_revision="revision-b",
            impact_run_id=run["id"],
            acknowledge_kcs_impact=True,
        )
    assert store.get_clause(project_id, "reviewed")["review_note"] == "영향 재검토 완료"


def test_quality_candidate_snapshot_is_unchanged_by_material_rematch(tmp_path):
    store = Store(tmp_path / "quality-invariant.sqlite3")
    project_id = _create_project(store, tmp_path)
    evaluation = store.ensure_quality_evaluation(project_id)
    assert evaluation
    store.update_quality_item(
        project_id,
        "reviewed",
        "candidate_selected",
        "candidate-1",
        "",
        "",
        "",
    )
    before_export = store.quality_export_rows(project_id)
    run, prepared = _queue_and_claim(store)
    results = _results(prepared)
    reviewed = next(item for item in results if item["id"] == "reviewed")
    reviewed["candidates"][0]["content"] = "품질평가 뒤 변경된 KCS 의미"

    store.apply_kcs_rematch(run["id"], prepared["expected_state_sha256"], results)

    after = store.get_quality_evaluation(project_id)
    after_item = next(item for item in after["items"] if item["clause_id"] == "reviewed")
    after_export = store.quality_export_rows(project_id)
    assert after_item["candidate_count"] == 2
    assert after_item["relevant_rank"] == 1
    assert after["metrics"]["rank_hits"]["1"] == 1
    assert after_export[0]["candidates"] == before_export[0]["candidates"]


def test_public_run_counts_unacknowledged_impacts_and_reason_is_readable(tmp_path):
    store = Store(tmp_path / "public-impact.sqlite3")
    project_id = _create_project(store, tmp_path, ready_decision="keep")
    run, prepared = _queue_and_claim(store)
    results = _results(prepared)
    reviewed = next(item for item in results if item["id"] == "reviewed")
    reviewed["candidates"][0]["content"] = "감사 대상이 되는 개정 KCS 본문"

    completed = store.apply_kcs_rematch(
        run["id"], prepared["expected_state_sha256"], results
    )

    assert completed["unacknowledged_count"] == 1
    impact = completed["impacts"][0]
    assert "후보 KCS의 의미가 변경됨" in impact["reason"]
    assert "candidate_meaning_changed" in impact["reasons"]
    assert "candidate_meaning_changed" not in impact["reason"]
    assert store.list_kcs_rematch_runs(project_id)[0]["unacknowledged_count"] == 1

    store.acknowledge_kcs_impact(
        project_id,
        "reviewed",
        run["id"],
        "revision-b",
    )
    assert store.get_kcs_rematch_run(run["id"])["unacknowledged_count"] == 0


def test_kcs_rematch_audit_rows_and_workbook_preserve_snapshots(tmp_path):
    store = Store(tmp_path / "audit-impact.sqlite3")
    project_id = _create_project(store, tmp_path, ready_decision="keep")
    run, prepared = _queue_and_claim(store)
    results = _results(prepared)
    reviewed = next(item for item in results if item["id"] == "reviewed")
    reviewed["candidates"][0]["content"] = "감사 엑셀에 남을 개정 KCS 본문"
    store.apply_kcs_rematch(run["id"], prepared["expected_state_sha256"], results)

    audit_rows = store.kcs_rematch_audit_rows(project_id)

    assert len(audit_rows) == 1
    audit = audit_rows[0]
    assert audit["run_id"] == run["id"]
    assert audit["status"] == "completed"
    assert audit["matcher_signature"] == {"mode": "test"}
    assert audit["clause_id"] == "reviewed"
    assert audit["before_candidates"][0]["content"] == "KCS 의미 본문 1"
    assert audit["after_candidates"][0]["content"] == "감사 엑셀에 남을 개정 KCS 본문"
    assert audit["before_decision"]["decision"] == "keep"

    project = store.get_project(project_id)
    workbook_bytes = build_audit_xlsx(project, store.audit_rows(project_id), audit_rows)
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True)

    assert workbook.sheetnames == ["판정이력", "KCS 개정영향", "승인이력"]
    sheet = workbook["KCS 개정영향"]
    headers = {cell.value: index for index, cell in enumerate(sheet[2], start=1)}
    assert sheet.cell(3, headers["작업 상태"]).value == "완료"
    assert "후보 KCS의 의미가 변경됨" in sheet.cell(3, headers["변경 사유"]).value
    assert "KCS 의미 본문 1" in sheet.cell(3, headers["변경 전 후보 JSON"]).value
    assert "감사 엑셀에 남을 개정 KCS 본문" in sheet.cell(
        3, headers["변경 후 후보 JSON"]
    ).value

    compatible = load_workbook(BytesIO(build_audit_xlsx(project, [])), read_only=True)
    assert compatible.sheetnames == ["판정이력", "KCS 개정영향", "승인이력"]
    assert compatible["KCS 개정영향"].max_row == 2


def test_recover_interrupted_rematch_requeues_valid_run_and_clears_progress(tmp_path):
    store = Store(tmp_path / "recover-valid.sqlite3")
    _create_project(store, tmp_path)
    run, prepared = _queue_and_claim(store)
    with store.connect() as connection:
        connection.execute(
            """
            UPDATE kcs_rematch_runs
            SET matched_count = 1, material_change_count = 1,
                review_required_count = 1, error = '중간 상태'
            WHERE id = ?
            """,
            (run["id"],),
        )
        connection.execute(
            """
            INSERT INTO kcs_clause_impacts(
                run_id, clause_id, impact_type, review_required,
                reason_json, before_candidates_json,
                after_candidates_json, before_decision_json
            ) VALUES (?, 'reviewed', 'material_change', 1,
                      '["candidate_removed"]', '[]', '[]', '{}')
            """,
            (run["id"],),
        )

    recovered = store.recover_interrupted_kcs_rematches()

    assert recovered["superseded"] == []
    assert [item["id"] for item in recovered["requeued"]] == [run["id"]]
    current = store.get_kcs_rematch_run(run["id"])
    assert current["status"] == "pending"
    assert current["expected_state_sha256"] == ""
    assert current["matcher_signature"] == {}
    assert current["total_clauses"] == 0
    assert current["matched_count"] == 0
    assert current["material_change_count"] == 0
    assert current["review_required_count"] == 0
    assert current["unacknowledged_count"] == 0
    assert current["error"] == ""
    assert current["started_at"] is None
    assert current["finished_at"] is None
    assert current["impacts"] == []
    assert store.recover_interrupted_kcs_rematches() == {
        "requeued": [],
        "superseded": [],
    }
    reclaimed = store.claim_kcs_rematch(run["id"], {"mode": "recovered"})
    assert reclaimed["expected_state_sha256"] == prepared["expected_state_sha256"]


@pytest.mark.parametrize("invalid_state", ["target", "parser", "from_revision"])
def test_recover_interrupted_rematch_supersedes_invalid_run(tmp_path, invalid_state):
    store = Store(tmp_path / f"recover-invalid-{invalid_state}.sqlite3")
    project_id = _create_project(store, tmp_path)
    run, _prepared = _queue_and_claim(store)
    with store.connect() as connection:
        if invalid_state == "target":
            connection.execute(
                "UPDATE kcs_rematch_runs SET target_snapshot = 'obsolete' WHERE id = ?",
                (run["id"],),
            )
        elif invalid_state == "parser":
            connection.execute(
                "UPDATE projects SET parser_version = '' WHERE id = ?",
                (project_id,),
            )
        else:
            connection.execute(
                "UPDATE projects SET kcs_revision = 'revision-diverged' WHERE id = ?",
                (project_id,),
            )
        connection.execute(
            """
            UPDATE kcs_rematch_runs
            SET matched_count = 1, material_change_count = 1,
                review_required_count = 1
            WHERE id = ?
            """,
            (run["id"],),
        )
        connection.execute(
            """
            INSERT INTO kcs_clause_impacts(
                run_id, clause_id, impact_type, review_required,
                reason_json, before_candidates_json,
                after_candidates_json, before_decision_json
            ) VALUES (?, 'reviewed', 'material_change', 1,
                      '["candidate_removed"]', '[]', '[]', '{}')
            """,
            (run["id"],),
        )

    recovered = store.recover_interrupted_kcs_rematches()

    assert recovered["requeued"] == []
    assert [item["id"] for item in recovered["superseded"]] == [run["id"]]
    current = store.get_kcs_rematch_run(run["id"])
    assert current["status"] == "superseded"
    assert current["expected_state_sha256"] == ""
    assert current["matcher_signature"] == {}
    assert current["total_clauses"] == 0
    assert current["matched_count"] == 0
    assert current["material_change_count"] == 0
    assert current["review_required_count"] == 0
    assert current["unacknowledged_count"] == 0
    assert current["error"]
    assert current["finished_at"] is not None
    assert current["impacts"] == []
