from __future__ import annotations

import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import uuid
import zipfile
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .text_analysis import starts_with_quantity


IS_WINDOWS = os.name == "nt"


def legacy_doc_conversion_engine() -> str | None:
    """Return the conversion engine available to the current runtime."""
    if IS_WINDOWS:
        return "microsoft_word"
    if shutil.which("libreoffice") or shutil.which("soffice"):
        return "libreoffice"
    return None


AI_RELATION_LABELS = {
    "equivalent": "실질적 동일",
    "kcs_covers": "KCS가 포괄",
    "partial_overlap": "부분 일치",
    "posco_specific": "포스코 특화 포함",
    "conflict": "기준 충돌",
    "unrelated": "무관 후보",
}

QUALITY_VERDICT_LABELS = {
    "candidate_selected": "제시 후보 적정",
    "all_candidates_incorrect": "적용 KCS 없음(후보 오탐)",
    "no_candidate_correct": "후보 없음이 적정",
    "kcs_missing": "정답 KCS Top3 누락(미탐)",
}

QUALITY_COHORT_LABELS = {
    "high": "관련성 높음 표본",
    "review": "검토 필요 표본",
    "no_candidate": "후보 없음 표본",
}

QUALITY_SCORE_BAND_LABELS = {
    "gte_0_45": "0.45 이상",
    "0_35_to_0_45": "0.35 이상 ~ 0.45 미만",
    "0_25_to_0_35": "0.25 이상 ~ 0.35 미만",
    "no_candidate": "후보 없음",
}

QUALITY_STATUS_LABELS = {
    "collecting": "평가 전",
    "partial": "평가 진행 중",
    "complete": "평가 완료",
}

KCS_REMATCH_STATUS_LABELS = {
    "pending": "대기",
    "running": "진행 중",
    "completed": "완료",
    "failed": "실패",
    "superseded": "대체됨",
}

KCS_IMPACT_TYPE_LABELS = {
    "material_change": "의미 변경",
    "metadata_change": "순위·점수·개정정보 변경",
}


LABEL_RE = re.compile(
    r"^\s*(?P<label>(?:\d+(?:\.\d+)*(?:\([^)]+\))?|[가나다라마바사아자차카타파하][.)]|[①-⑳]|\([0-9가나다라마바사아자차카타파하]+\)))\s*[.)]?\s*(?P<body>.*)$"
)
SENTENCE_END_RE = re.compile(r"(?:다|함|됨|한다|된다|있다|없다)[.]?$|[.!?。]$")
TOC_PAGE_LOCATOR_RE = re.compile(r"(?:\t+|[.·ㆍ…]{2,})\s*\d+\s*$")
REQUIREMENT_SIGNAL_RE = re.compile(
    r"(?:하여야|해야|한다|된다|따른다|없어야|있어야|금지|허용|"
    r"이상|이하|초과|미만|이내|이후|이전|까지|마다|"
    r"사용한다|설치한다|시공한다|확인한다|검사한다|측정한다)"
)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def clean_text(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def safe_filename(value: str) -> str:
    name = re.sub(r'[<>:"/\\|?*]+', "_", value).strip(" .")
    return name[:120] or "시방서"


EXCEL_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _excel_safe(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    cleaned = EXCEL_CONTROL_RE.sub("", value)
    if cleaned.startswith(("=", "+", "-", "@")):
        return "'" + cleaned[:32766]
    return cleaned[:32767]


def _excel_row(values: Iterable[Any]) -> list[Any]:
    return [_excel_safe(value) for value in values]


def convert_legacy_doc(input_path: Path, output_path: Path) -> bytes:
    input_path = input_path.resolve()
    output_path = output_path.resolve()
    if input_path.suffix.lower() != ".doc" or output_path.suffix.lower() != ".docx":
        raise ValueError("DOC 변환 경로가 올바르지 않습니다.")
    if input_path.parent != output_path.parent:
        raise ValueError("DOC와 변환 파일은 같은 보안 저장 폴더에 있어야 합니다.")

    output_path.unlink(missing_ok=True)
    if IS_WINDOWS:
        script_path = Path(__file__).with_name("convert_legacy_doc.ps1")
        command = [
            "powershell.exe",
            "-NoLogo",
            "-NoProfile",
            "-NonInteractive",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script_path),
            "-InputPath",
            str(input_path),
            "-OutputPath",
            str(output_path),
        ]
    else:
        office = shutil.which("libreoffice") or shutil.which("soffice")
        if not office:
            raise ValueError("DOC 변환 프로그램(LibreOffice)을 찾을 수 없습니다.")
        command = [
            office,
            "--headless",
            "--convert-to",
            "docx",
            "--outdir",
            str(input_path.parent),
            str(input_path),
        ]
    try:
        result = subprocess.run(
            command,
            capture_output=True,
            timeout=180,
            check=False,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        output_path.unlink(missing_ok=True)
        raise ValueError("DOC 파일 변환 프로그램을 시작하지 못했습니다.") from exc
    if result.returncode != 0 or not output_path.is_file() or output_path.stat().st_size == 0:
        output_path.unlink(missing_ok=True)
        raise ValueError("DOC 파일을 변환하지 못했습니다. 암호 또는 파일 손상 여부를 확인하세요.")
    try:
        with zipfile.ZipFile(output_path) as archive:
            if "word/document.xml" not in archive.namelist():
                raise ValueError
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        output_path.unlink(missing_ok=True)
        raise ValueError("변환 결과가 올바른 DOCX 파일이 아닙니다.") from exc
    return output_path.read_bytes()


def _iter_blocks(document: DocumentObject) -> Iterable[Paragraph | Table]:
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


class NumberingResolver:
    """Resolve common Word list numbering stored directly on a paragraph/style."""

    def __init__(self, document: DocumentObject):
        self.document = document
        self.counters: dict[str, list[int]] = {}
        self.numbering = None
        try:
            self.numbering = document.part.numbering_part.element
        except (AttributeError, KeyError, NotImplementedError):
            pass

    def _num_pr(self, paragraph: Paragraph):
        p_pr = paragraph._p.pPr
        if p_pr is not None and p_pr.numPr is not None:
            return p_pr.numPr
        try:
            style_p_pr = paragraph.style._element.pPr
            if style_p_pr is not None and style_p_pr.numPr is not None:
                return style_p_pr.numPr
        except AttributeError:
            pass
        return None

    def label(self, paragraph: Paragraph) -> str:
        num_pr = self._num_pr(paragraph)
        if num_pr is None or num_pr.numId is None:
            return ""
        num_id = str(num_pr.numId.val)
        level = int(num_pr.ilvl.val) if num_pr.ilvl is not None else 0
        if self.numbering is None:
            return ""

        num_nodes = self.numbering.xpath(f'./w:num[@w:numId="{num_id}"]')
        if not num_nodes:
            return ""
        abstract_id_node = num_nodes[0].find(qn("w:abstractNumId"))
        if abstract_id_node is None:
            return ""
        abstract_id = abstract_id_node.get(qn("w:val"))
        level_nodes = self.numbering.xpath(
            f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]/w:lvl[@w:ilvl="{level}"]'
        )
        if not level_nodes:
            return ""
        level_node = level_nodes[0]
        fmt_node = level_node.find(qn("w:numFmt"))
        text_node = level_node.find(qn("w:lvlText"))
        start_node = level_node.find(qn("w:start"))
        number_format = fmt_node.get(qn("w:val")) if fmt_node is not None else "decimal"
        pattern = text_node.get(qn("w:val")) if text_node is not None else f"%{level + 1}."
        start = int(start_node.get(qn("w:val"))) if start_node is not None else 1

        counters = self.counters.setdefault(num_id, [0] * 9)
        if counters[level] == 0:
            counters[level] = start
        else:
            counters[level] += 1
        for deeper in range(level + 1, len(counters)):
            counters[deeper] = 0

        if number_format == "bullet":
            return "•"

        def render_number(index: int, value: int) -> str:
            if value <= 0:
                value = 1
            if number_format == "lowerLetter" and index == level:
                return chr(ord("a") + ((value - 1) % 26))
            if number_format == "upperLetter" and index == level:
                return chr(ord("A") + ((value - 1) % 26))
            return str(value)

        rendered = pattern
        for index in range(9):
            rendered = rendered.replace(f"%{index + 1}", render_number(index, counters[index]))
        return rendered


def _outline_level(paragraph: Paragraph) -> int | None:
    style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
    heading = re.search(r"(?:heading|제목)\s*(\d+)", style_name)
    if heading:
        return min(9, max(1, int(heading.group(1))))
    try:
        outline = paragraph._p.pPr.outlineLvl
        if outline is not None:
            return int(outline.val) + 1
    except AttributeError:
        pass
    return None


def _split_label(text: str) -> tuple[str, str]:
    match = LABEL_RE.match(text)
    if not match:
        return "", text
    return clean_text(match.group("label")), clean_text(match.group("body"))


def _title_and_content(text: str, is_heading: bool) -> tuple[str, str]:
    colon = re.match(r"^(?P<title>.{2,60}?)\s*[:：]\s*(?P<content>\S.*)$", text)
    if colon:
        return clean_text(colon.group("title")), clean_text(colon.group("content"))
    if is_heading or (len(text) <= 90 and not SENTENCE_END_RE.search(text)):
        return text, ""
    title = text[:90].rstrip()
    if len(text) > 90:
        title += "…"
    return title, text


def _looks_like_numbered_heading(label: str, body: str) -> bool:
    """Detect short numbered section names without hiding actual requirements."""
    if not label or not body or len(body) > 50:
        return False
    if SENTENCE_END_RE.search(body) or REQUIREMENT_SIGNAL_RE.search(body):
        return False
    return bool(re.search(r"[가-힣A-Za-z]", body))


def _normalise_toc_text(value: str) -> str:
    return re.sub(r"[^0-9a-z가-힣]", "", clean_text(value).casefold())


def _has_explicit_numeric_label_evidence(
    raw_text: str,
    label: str,
    body: str,
    toc_titles: dict[str, str],
) -> bool:
    if not re.fullmatch(r"\d+(?:\.\d+)*", label):
        return True
    has_delimiter = bool(
        re.match(rf"^\s*{re.escape(label)}\s*[.)]", raw_text)
    )
    toc_title = toc_titles.get(label)
    normalized_body = _normalise_toc_text(body)
    normalized_toc_title = _normalise_toc_text(toc_title or "")
    exact_toc_match = bool(
        normalized_toc_title and normalized_body == normalized_toc_title
    )
    near_toc_match = bool(
        toc_title
        and len(clean_text(body)) <= 30
        and len(clean_text(toc_title)) <= 30
        and not SENTENCE_END_RE.search(body)
        and not SENTENCE_END_RE.search(toc_title)
        and abs(len(normalized_body) - len(normalized_toc_title)) <= 1
        and (
            normalized_body in normalized_toc_title
            or normalized_toc_title in normalized_body
        )
    )
    matches_toc = exact_toc_match or near_toc_match
    bare_integer = bool(re.fullmatch(r"\d+", label))
    zero_leading_number = label.split(".", 1)[0] == "0"
    leading_quantity = starts_with_quantity(raw_text)
    return not (
        bare_integer or zero_leading_number or leading_quantity
    ) or has_delimiter or matches_toc


def _has_toc_metadata(paragraph: Paragraph) -> bool:
    style_name = (paragraph.style.name or "").casefold() if paragraph.style else ""
    if re.search(r"(?:^|\s)(?:toc|목차)\s*\d*", style_name):
        return True
    field_code = " ".join(node.text or "" for node in paragraph._p.xpath(".//w:instrText"))
    return "TOC" in field_code.upper()


def _has_page_break(paragraph: Paragraph) -> bool:
    return bool(
        paragraph._p.xpath('.//w:br[@w:type="page"] | .//w:lastRenderedPageBreak')
    )


def _has_toc_page_locator(paragraph: Paragraph) -> bool:
    return bool(TOC_PAGE_LOCATOR_RE.search(paragraph.text or ""))


def _toc_entry(paragraph: Paragraph) -> tuple[str, str] | None:
    raw_text = paragraph.text or ""
    raw_text = TOC_PAGE_LOCATOR_RE.sub("", raw_text).strip()
    label, title = _split_label(clean_text(raw_text))
    if not re.fullmatch(r"\d+(?:\.\d+)*", label) or not title or len(title) > 200:
        return None
    return label, title


def _same_toc_entry(left: tuple[str, str], right: tuple[str, str]) -> bool:
    return left[0] == right[0] and _normalise_toc_text(left[1]) == _normalise_toc_text(right[1])


def _toc_result(
    blocks: list[Paragraph | Table],
    body_start: int,
    entries: list[tuple[str, str]],
) -> tuple[int, int, dict[str, str]]:
    titles: dict[str, str] = {}
    for label, title in entries:
        titles.setdefault(label, title)
    skipped = sum(
        1
        for block in blocks[:body_start]
        if isinstance(block, Paragraph) and clean_text(block.text)
    )
    return body_start, skipped, titles


def _next_content_block(blocks: list[Paragraph | Table], start: int) -> int | None:
    for index in range(start, len(blocks)):
        block = blocks[index]
        if isinstance(block, Table) or clean_text(block.text):
            return index
    return None


def _matches_toc_anchor(
    blocks: list[Paragraph | Table],
    start: int,
    expected: list[tuple[str, str]],
) -> bool:
    """Require three ordered numeric anchors before inferring an implicit TOC."""

    if len(expected) < 3:
        return False
    matched = 0
    for block in blocks[start:]:
        if not isinstance(block, Paragraph):
            continue
        entry = _toc_entry(block)
        if entry is None:
            continue
        if not _same_toc_entry(expected[matched], entry):
            return False
        matched += 1
        if matched == 3:
            return True
    return False


def _detect_toc_region(
    blocks: list[Paragraph | Table],
) -> tuple[int, int, dict[str, str]]:
    """Find a high-confidence front-matter TOC without dropping uncertain body text."""

    scan_limit = min(len(blocks), 200)
    marker_index: int | None = None
    entry_start = 0
    for index, block in enumerate(blocks[:scan_limit]):
        if not isinstance(block, Paragraph):
            continue
        marker = _normalise_toc_text(block.text) in {
            "목차",
            "차례",
            "contents",
            "tableofcontents",
        }
        if marker or _has_toc_metadata(block):
            marker_index = index
            entry_start = index + 1 if marker else index
            break

    if marker_index is not None:
        entries: list[tuple[str, str]] = []
        first_entry: tuple[str, str] | None = None
        search_end = min(len(blocks), marker_index + 400)
        for index in range(entry_start, search_end):
            block = blocks[index]
            if not isinstance(block, Paragraph):
                continue
            entry = _toc_entry(block)
            if entry:
                if first_entry is None:
                    first_entry = entry
                elif len({label for label, _ in entries}) >= 2 and _same_toc_entry(
                    first_entry, entry
                ):
                    return _toc_result(blocks, index, entries)
                entries.append(entry)
            if _has_page_break(block) and len({label for label, _ in entries}) >= 2:
                body_start = _next_content_block(blocks, index + 1)
                if body_start is not None:
                    next_block = blocks[body_start]
                    if (
                        isinstance(next_block, Paragraph)
                        and _toc_entry(next_block)
                        and (
                            _has_toc_metadata(next_block)
                            or _has_toc_page_locator(next_block)
                        )
                    ):
                        continue
                    return _toc_result(blocks, body_start, entries)

    run: list[tuple[int, tuple[str, str]]] = []
    completed_runs: list[list[tuple[int, tuple[str, str]]]] = []
    for index, block in enumerate(blocks[:scan_limit]):
        if not isinstance(block, Paragraph) or not clean_text(block.text):
            continue
        entry = _toc_entry(block)
        if not entry:
            if len(run) >= 5:
                completed_runs.append(run)
            run = []
            continue
        if (
            len(run) >= 5
            and "." not in run[0][1][0]
            and _same_toc_entry(run[0][1], entry)
            and _matches_toc_anchor(
                blocks,
                index,
                [item for _, item in run[:3]],
            )
        ):
            return _toc_result(blocks, index, [item for _, item in run])
        run.append((index, entry))
    if len(run) >= 5:
        completed_runs.append(run)

    for candidate_run in completed_runs:
        _, first_entry = candidate_run[0]
        if "." in first_entry[0]:
            continue
        for index in range(candidate_run[-1][0] + 1, len(blocks)):
            block = blocks[index]
            if not isinstance(block, Paragraph):
                continue
            entry = _toc_entry(block)
            if (
                entry
                and _same_toc_entry(first_entry, entry)
                and _matches_toc_anchor(
                    blocks,
                    index,
                    [item for _, item in candidate_run[:3]],
                )
            ):
                return _toc_result(
                    blocks,
                    index,
                    [item for _, item in candidate_run],
                )
    return 0, 0, {}


def _contextual_label(
    label: str,
    numeric_context: str,
    parent_context: str,
) -> tuple[str, str, str]:
    compact = clean_text(label).replace(" ", "")
    numeric = re.fullmatch(r"(?P<base>\d+(?:\.\d+)*)(?P<child>\([^)]+\))?", compact)
    if numeric:
        base = numeric.group("base")
        child = numeric.group("child") or ""
        full = f"{base}{child}"
        return full, base, full if child else ""
    if re.fullmatch(r"\([0-9가나다라마바사아자차카타파하]+\)", compact):
        full = f"{numeric_context}{compact}" if numeric_context else compact
        return full, numeric_context, full
    if re.fullmatch(r"[가나다라마바사아자차카타파하][.)]", compact):
        context = parent_context or numeric_context
        full = f"{context}-{compact}" if context else compact
        return full, numeric_context, parent_context
    return compact, numeric_context, parent_context


def _section_context(
    label: str,
    parent_context: str,
    section_titles: dict[str, str],
    include_current: bool = False,
) -> str:
    numeric = re.match(r"\d+(?:\.\d+)*", label)
    if not numeric:
        return ""
    base = numeric.group(0)
    parts = base.split(".")
    keys = [".".join(parts[:index]) for index in range(1, len(parts) + 1)]
    if not include_current and label == base:
        keys = keys[:-1]
    if parent_context and parent_context != label:
        keys.append(parent_context)
    if include_current:
        keys.append(label)
    rendered: list[str] = []
    seen: set[str] = set()
    for key in keys:
        if key in seen or key not in section_titles:
            continue
        seen.add(key)
        rendered.append(clean_text(f"{key} {section_titles[key]}"))
    return " > ".join(rendered)


def _append_continuation(clause: dict[str, Any], text: str) -> None:
    if clause["content"]:
        clause["content"] = clean_text(f"{clause['content']} {text}")
    else:
        clause["content"] = text


def parse_docx(data: bytes, filename: str, project_id: str) -> tuple[str, list[dict[str, Any]], list[str]]:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("DOCX 파일을 열 수 없습니다. 손상되었거나 암호화된 파일인지 확인하세요.") from exc

    blocks = list(_iter_blocks(document))
    body_start, skipped_toc_items, toc_titles = _detect_toc_region(blocks)
    resolver = NumberingResolver(document)
    clauses: list[dict[str, Any]] = []
    warnings: list[str] = []
    source_order = 0
    table_count = 0
    automatic_numbering_count = 0
    merged_continuation_count = 0
    numeric_context = ""
    parent_context = ""
    section_titles = dict(toc_titles)
    current_match_context = ""

    for block in blocks[body_start:]:
        if isinstance(block, Paragraph):
            raw_text = clean_text(block.text)
            if not raw_text:
                continue
            generated_label = resolver.label(block)
            if generated_label:
                automatic_numbering_count += 1
            explicit_label, body = _split_label(raw_text)
            if generated_label:
                explicit_label = ""
                body = raw_text
            elif (
                explicit_label
                and not _has_explicit_numeric_label_evidence(
                    raw_text,
                    explicit_label,
                    body,
                    toc_titles,
                )
            ):
                explicit_label = ""
                body = raw_text
            detected_label = generated_label or explicit_label
            if not detected_label and clauses and clauses[-1]["source_type"] != "table":
                _append_continuation(clauses[-1], raw_text)
                if (
                    clauses[-1]["source_type"] == "heading"
                    and SENTENCE_END_RE.search(raw_text)
                ):
                    clauses[-1]["source_type"] = "paragraph"
                merged_continuation_count += 1
                continue

            label = detected_label or f"문단 {source_order + 1}"
            if detected_label:
                label, numeric_context, parent_context = _contextual_label(
                    label, numeric_context, parent_context
                )
            body = body if explicit_label else raw_text
            outline_level = _outline_level(block)
            is_top_heading = (
                bool(re.fullmatch(r"\d+", label))
                and len(body) <= 90
                and not SENTENCE_END_RE.search(body)
            )
            is_heading = (
                outline_level is not None
                or is_top_heading
                or (bool(detected_label) and _looks_like_numbered_heading(label, body))
            )
            title, content = _title_and_content(body, is_heading)
            match_context = (
                _section_context(label, parent_context, section_titles)
                if detected_label
                else current_match_context
            )
            source_order += 1
            clause_id = str(uuid.uuid5(uuid.UUID(project_id), f"p:{source_order}:{raw_text[:160]}"))
            clauses.append(
                {
                    "id": clause_id,
                    "source_order": source_order,
                    "label": label,
                    "title": title,
                    "content": content,
                    "source_type": "heading" if is_heading else "paragraph",
                    "outline_level": outline_level,
                    "match_context": match_context,
                }
            )
            if detected_label and title and not content:
                section_titles[label] = title
                current_match_context = _section_context(
                    label,
                    parent_context,
                    section_titles,
                    include_current=True,
                )
            continue

        table_count += 1
        for row_index, row in enumerate(block.rows, start=1):
            cell_texts = [clean_text(cell.text) for cell in row.cells]
            cell_texts = list(dict.fromkeys(text for text in cell_texts if text))
            if not cell_texts:
                continue
            source_order += 1
            joined = " | ".join(cell_texts)
            clause_id = str(uuid.uuid5(uuid.UUID(project_id), f"t:{source_order}:{joined[:160]}"))
            clauses.append(
                {
                    "id": clause_id,
                    "source_order": source_order,
                    "label": f"표 {table_count}-{row_index}",
                    "title": cell_texts[0][:90],
                    "content": joined,
                    "source_type": "table",
                    "outline_level": None,
                    "match_context": current_match_context,
                }
            )

    if not clauses:
        raise ValueError("검토할 문장이나 표 내용을 찾지 못했습니다.")
    if skipped_toc_items:
        warnings.append(f"목차와 표제부 {skipped_toc_items}개 항목을 검토 대상에서 제외했습니다.")
    if merged_continuation_count:
        warnings.append(f"줄바꿈된 문단 {merged_continuation_count}개를 앞 조항에 연결했습니다.")
    if table_count:
        warnings.append(f"표 {table_count}개를 행 단위 검토 조항으로 변환했습니다.")
    if automatic_numbering_count:
        warnings.append(f"Word 자동번호 문단 {automatic_numbering_count}개를 복원했습니다. 번호를 화면에서 확인하세요.")
    if (
        next(document.element.body.iter(qn("w:ins")), None) is not None
        or next(document.element.body.iter(qn("w:del")), None) is not None
    ):
        warnings.append("변경 내용 추적 요소가 있습니다. 확정된 문서로 다시 업로드하는 것을 권장합니다.")

    core_title = clean_text(document.core_properties.title)
    title = (
        Path(filename).stem
        if core_title.casefold() in {"word document", "document", "문서"}
        else core_title or Path(filename).stem
    )
    return title, clauses, warnings


def build_review_docx(project: dict[str, Any], clauses: list[dict[str, Any]], kind: str) -> bytes:
    document = Document()
    heading = document.add_heading(project["title"], level=0)
    heading.style = document.styles["Title"]
    subtitle = document.add_paragraph()
    subtitle.add_run("검토용 간소화 시방서" if kind == "review" else "KCS 보완·특기 시방서").bold = True
    if kind == "review":
        document.add_paragraph(f"원본: {project['source_filename']}")
        document.add_paragraph(f"비교 KCS 스냅샷: {project['kcs_snapshot']} / 범위: {project['kcs_scope']}")
        notice = document.add_paragraph("보류 조항은 최종 승인 전까지 현장 적용 기준으로 사용할 수 없습니다.")
        notice.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
    else:
        document.add_paragraph(f"적용 KCS 기준: {project['kcs_snapshot']}")
        document.add_paragraph(
            "KCS와 중복되지 않아 유지하기로 확정한 포스코 추가·강화 기준만 수록합니다."
        )

    document.add_page_break()
    for clause in clauses:
        label_title = clean_text(f"{clause['label']} {clause['title']}")
        level = min(3, max(1, clause.get("outline_level") or 2))
        clause_heading = document.add_heading(label_title, level=level)
        if clause["decision"] == "hold":
            prefix = clause_heading.insert_paragraph_before("[보류]")
            prefix.runs[0].bold = True
            prefix.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW

        edited_content = clean_text(clause.get("edited_content"))
        original_content = clean_text(clause.get("content") or clause.get("title"))
        if kind == "final" and clause.get("decision_reason") == "partial_overlap_residual":
            if not edited_content or edited_content == original_content:
                raise ValueError(
                    "부분 중복 잔여기준의 수정문이 없어 최종 문서를 생성할 수 없습니다."
                )
        final_content = edited_content or original_content
        if final_content and final_content != clause["title"]:
            document.add_paragraph(final_content)
        if kind == "review" and clause.get("kcs_code"):
            reference = document.add_paragraph(
                f"연결 KCS: {clause['kcs_code']} {clause.get('kcs_clause') or ''} · {clause.get('kcs_document_name') or ''}"
            )
            for run in reference.runs:
                run.italic = True
        if kind == "review" and clause.get("review_note"):
            note = document.add_paragraph(f"검토의견: {clause['review_note']}")
            note.style = document.styles["Quote"]

    if kind == "final" and not clauses:
        document.add_paragraph("검토 결과 별도로 유지할 포스코 추가·강화 조항이 없습니다.")

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def build_audit_xlsx(
    project: dict[str, Any],
    rows: list[dict[str, Any]],
    kcs_rematch_rows: list[dict[str, Any]] | None = None,
    workflow_rows: list[dict[str, Any]] | None = None,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "판정이력"
    headers = [
        "순서", "포스코 조항", "포스코 제목", "포스코 원문", "현재 판정",
        "판정사유", "전체포괄 확인", "수정문", "검토의견", "검토일시", "선택 KCS", "KCS 문서명",
        "KCS 조항", "유사도", "GPT 관계", "GPT 신뢰도", "GPT 판정근거",
        "GPT 간소화 제안", "이전 판정", "변경 판정", "변경일시",
        "전체포괄 상태", "전체포괄 신뢰도", "전체포괄 삭제안전",
        "전체포괄 근거 KCS", "요구사항별 분석", "잔여 포스코 문구", "전체포괄 판정근거",
        "전체포괄 모델", "전체포괄 분석일시", "요구사항 원본 JSON",
    ]
    sheet.append(_excel_row([project["title"], project["source_filename"], project["kcs_snapshot"]]))
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    sheet.append(_excel_row(headers))
    for row in rows:
        sheet.append(
            _excel_row([
                row.get("source_order"), row.get("label"), row.get("title"), row.get("content"),
                row.get("decision"), row.get("decision_reason"), "예" if row.get("coverage_confirmed") else "아니오",
                row.get("edited_content"), row.get("review_note"), row.get("reviewed_at"),
                row.get("kcs_code"), row.get("kcs_document_name"), row.get("kcs_clause"),
                row.get("score"), AI_RELATION_LABELS.get(row.get("ai_relation_type"), row.get("ai_relation_type")), row.get("ai_confidence"),
                row.get("ai_rationale"), row.get("ai_simplified_content"),
                row.get("previous_decision"), row.get("new_decision"), row.get("changed_at"),
                row.get("coverage_status"), row.get("coverage_confidence"),
                "예" if row.get("coverage_deletion_safe") else "아니오" if row.get("coverage_deletion_safe") is not None else "",
                row.get("coverage_evidence_kcs"), row.get("coverage_requirements"),
                row.get("coverage_residual_content"), row.get("coverage_rationale"),
                row.get("coverage_model"), row.get("coverage_analyzed_at"),
                row.get("coverage_requirements_json"),
            ])
        )

    navy = "17365D"
    sheet["A1"].fill = PatternFill("solid", fgColor=navy)
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    for cell in sheet[2]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{max(2, sheet.max_row)}"
    widths = [8, 14, 28, 55, 12, 24, 14, 55, 35, 22, 18, 28, 14, 10, 18, 12, 45, 55,
              12, 12, 22, 18, 14, 16, 45, 65, 45, 55, 18, 22, 65]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    impact_sheet = workbook.create_sheet("KCS 개정영향")
    impact_headers = [
        "작업 ID", "작업 상태", "시작 KCS 개정", "대상 KCS 개정", "대상 KCS 스냅샷",
        "대기 일시", "시작 일시", "완료 일시", "전체 매칭 대상", "후보 보유 조항",
        "의미 변경 수", "재검토 필요 수", "미확인 수", "오류", "매처 설정 JSON",
        "작업 상태 지문", "포스코 순서", "포스코 조항", "포스코 제목", "포스코 원문",
        "영향 유형", "재검토 필요", "확인 일시", "변경 사유", "변경 사유 코드",
        "변경 전 후보 요약", "변경 후 후보 요약", "변경 전 판정", "변경 전 판정사유",
        "변경 전 선택 후보", "변경 전 전체포괄 확인", "변경 전 후보 JSON",
        "변경 후 후보 JSON", "변경 전 판정 JSON",
    ]
    impact_sheet.append(
        _excel_row(
            [
                f"{project.get('title') or ''} · KCS 개정영향 감사이력",
                project.get("source_filename"),
                project.get("kcs_snapshot"),
            ]
        )
    )
    impact_sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(impact_headers),
    )
    impact_sheet.append(_excel_row(impact_headers))

    def audit_json(value: Any) -> str:
        if value in (None, "", [], {}):
            return ""
        return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)

    def candidate_summary(candidates: Any) -> str:
        if not isinstance(candidates, list):
            return ""
        lines: list[str] = []
        for candidate in candidates:
            if not isinstance(candidate, dict):
                continue
            reference = " ".join(
                str(candidate.get(key) or "").strip()
                for key in ("kcs_code", "kcs_clause")
                if candidate.get(key)
            )
            score = candidate.get("score")
            score_text = f" ({float(score):.1%})" if isinstance(score, (int, float)) else ""
            title = str(candidate.get("title") or "").strip()
            rank = candidate.get("rank")
            lines.append(
                f"{rank if rank is not None else '-'}순위 · {reference}{score_text} · {title}".strip()
            )
        return "\n".join(lines)

    for row in kcs_rematch_rows or []:
        before_decision = row.get("before_decision")
        if not isinstance(before_decision, dict):
            before_decision = {}
        reasons = row.get("reasons")
        reason_codes = ", ".join(str(reason) for reason in reasons) if isinstance(reasons, list) else ""
        impact_sheet.append(
            _excel_row(
                [
                    row.get("run_id"),
                    KCS_REMATCH_STATUS_LABELS.get(row.get("status"), row.get("status")),
                    row.get("from_revision"),
                    row.get("target_revision"),
                    row.get("target_snapshot"),
                    row.get("queued_at"),
                    row.get("started_at"),
                    row.get("finished_at"),
                    row.get("total_clauses"),
                    row.get("matched_count"),
                    row.get("material_change_count"),
                    row.get("review_required_count"),
                    row.get("unacknowledged_count"),
                    row.get("error"),
                    audit_json(row.get("matcher_signature")),
                    row.get("expected_state_sha256"),
                    row.get("source_order"),
                    row.get("label"),
                    row.get("title"),
                    row.get("content"),
                    KCS_IMPACT_TYPE_LABELS.get(
                        row.get("impact_type"), row.get("impact_type")
                    ),
                    "예" if row.get("review_required") else "아니오",
                    row.get("acknowledged_at"),
                    row.get("reason"),
                    reason_codes,
                    candidate_summary(row.get("before_candidates")),
                    candidate_summary(row.get("after_candidates")),
                    before_decision.get("decision"),
                    before_decision.get("decision_reason"),
                    before_decision.get("selected_candidate_id"),
                    "예" if before_decision.get("coverage_confirmed") else "아니오",
                    audit_json(row.get("before_candidates")),
                    audit_json(row.get("after_candidates")),
                    audit_json(before_decision),
                ]
            )
        )

    impact_sheet["A1"].fill = PatternFill("solid", fgColor=navy)
    impact_sheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    for cell in impact_sheet[2]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    impact_sheet.freeze_panes = "A3"
    impact_sheet.auto_filter.ref = (
        f"A2:{get_column_letter(len(impact_headers))}{max(2, impact_sheet.max_row)}"
    )
    impact_widths = [
        38, 12, 18, 18, 28, 24, 24, 24, 14, 14, 14, 16, 12, 38, 35, 22,
        12, 16, 30, 55, 24, 14, 24, 42, 38, 55, 55, 14, 24, 38, 18, 65, 65, 65,
    ]
    for index, width in enumerate(impact_widths, start=1):
        impact_sheet.column_dimensions[get_column_letter(index)].width = width
    for row in impact_sheet.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    workflow_sheet = workbook.create_sheet("승인이력")
    workflow_headers = [
        "검토 회차", "제출 ID", "처리 유형", "이전 상태", "변경 상태",
        "처리자", "작성자", "승인자", "메모·반려사유", "처리 일시",
        "제출 일시", "결정 일시", "KCS 개정", "KCS 스냅샷",
        "원본 SHA-256", "검토본 SHA-256",
    ]
    workflow_sheet.append(
        _excel_row(
            [
                f"{project.get('title') or ''} · 작성자–승인자 처리 이력",
                project.get("source_filename"),
                project.get("kcs_snapshot"),
            ]
        )
    )
    workflow_sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(workflow_headers),
    )
    workflow_sheet.append(_excel_row(workflow_headers))

    workflow_event_labels = {
        "submitted": "검토 제출",
        "approved": "승인",
        "changes_requested": "반려",
        "superseded": "제출본 무효화",
    }
    workflow_status_labels = {
        "reviewing": "작성 중",
        "submitted": "승인 대기",
        "changes_requested": "반려·수정 중",
        "approved": "승인 완료",
        "superseded": "무효화",
    }
    for row in workflow_rows or []:
        workflow_sheet.append(
            _excel_row(
                [
                    row.get("revision_no"),
                    row.get("submission_id") or row.get("id"),
                    workflow_event_labels.get(
                        row.get("event_type"), row.get("event_type")
                    ),
                    workflow_status_labels.get(
                        row.get("from_status"), row.get("from_status")
                    ),
                    workflow_status_labels.get(
                        row.get("to_status"), row.get("to_status")
                    ),
                    row.get("actor_name"),
                    row.get("author_name"),
                    row.get("decided_by") or row.get("approver_name"),
                    row.get("note") or row.get("decision_note"),
                    row.get("occurred_at") or row.get("created_at"),
                    row.get("submitted_at"),
                    row.get("decided_at"),
                    row.get("kcs_revision"),
                    row.get("kcs_snapshot"),
                    row.get("source_sha256"),
                    row.get("review_snapshot_sha256") or row.get("state_sha256"),
                ]
            )
        )

    workflow_sheet["A1"].fill = PatternFill("solid", fgColor=navy)
    workflow_sheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    for cell in workflow_sheet[2]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    workflow_sheet.freeze_panes = "A3"
    workflow_sheet.auto_filter.ref = (
        f"A2:{get_column_letter(len(workflow_headers))}{max(2, workflow_sheet.max_row)}"
    )
    workflow_widths = [
        12, 38, 18, 18, 18, 18, 18, 18, 48, 24, 24, 24, 24, 28, 66, 66,
    ]
    for index, width in enumerate(workflow_widths, start=1):
        workflow_sheet.column_dimensions[get_column_letter(index)].width = width
    for row in workflow_sheet.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_quality_evaluation_xlsx(
    project: dict[str, Any],
    evaluation: dict[str, Any],
    rows: list[dict[str, Any]],
) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "요약"
    navy = "17365D"
    pale_blue = "D9EAF7"

    summary.append(_excel_row(["KCS 매칭 품질평가 요약"]))
    summary.merge_cells("A1:I1")
    summary["A1"].fill = PatternFill("solid", fgColor=navy)
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    summary["A1"].alignment = Alignment(horizontal="left")
    insights = evaluation["insights"]
    metrics = evaluation["metrics"]
    summary_rows = [
        ("프로젝트", project.get("title")),
        ("원본 파일", project.get("source_filename")),
        ("KCS 스냅샷", project.get("kcs_snapshot")),
        ("KCS 범위", project.get("kcs_scope")),
        ("평가 담당자", evaluation.get("reviewer_name")),
        ("평가 상태", QUALITY_STATUS_LABELS.get(insights.get("status"), insights.get("status"))),
        ("안내", insights.get("message")),
        ("집계 범위", insights.get("scope_notice")),
        ("표본 수", evaluation.get("actual_sample_size")),
        ("평가 완료", metrics.get("evaluated_count")),
        ("남은 표본", metrics.get("remaining_count")),
        ("Recall@3", metrics.get("recall_at_3")),
        ("MRR", metrics.get("mrr")),
        ("평가 완료 표본 내 정확도", metrics.get("accuracy")),
    ]
    for label, value in summary_rows:
        summary.append(_excel_row([label, value]))
        if label in {"Recall@3", "평가 완료 표본 내 정확도"}:
            summary.cell(summary.max_row, 2).number_format = "0.00%"
        elif label == "MRR":
            summary.cell(summary.max_row, 2).number_format = "0.000"

    summary.append([])
    summary.append(_excel_row(["판정 정의"]))
    definition_title_row = summary.max_row
    summary.merge_cells(
        start_row=definition_title_row, start_column=1, end_row=definition_title_row, end_column=9
    )
    summary.cell(definition_title_row, 1).fill = PatternFill("solid", fgColor=pale_blue)
    summary.cell(definition_title_row, 1).font = Font(bold=True, color=navy)
    summary.append(
        _excel_row(
            [
                "제시 후보 모두 부적정",
                "후보는 있으나 적용 가능한 KCS가 없음(오탐)",
            ]
        )
    )
    summary.append(
        _excel_row(
            [
                "정답 KCS 누락",
                "정답 KCS가 존재하지만 상위 3개 후보에 없음(미탐)",
            ]
        )
    )

    summary.append([])
    summary.append(_excel_row(["판정 수"]))
    verdict_title_row = summary.max_row
    summary.merge_cells(
        start_row=verdict_title_row, start_column=1, end_row=verdict_title_row, end_column=2
    )
    summary.cell(verdict_title_row, 1).fill = PatternFill("solid", fgColor=pale_blue)
    summary.cell(verdict_title_row, 1).font = Font(bold=True, color=navy)
    summary.append(_excel_row(["판정", "건수"]))
    verdict_header_row = summary.max_row
    for cell in summary[verdict_header_row]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for verdict, label in QUALITY_VERDICT_LABELS.items():
        summary.append(_excel_row([label, metrics["counts"][verdict]]))

    verdict_headers = [
        "구분", "표본 수", "평가 수", "적정 수", "평가 표본 내 정확도", "제시 후보 적정",
        "후보 모두 부적정", "후보 없음 적정", "정답 KCS 누락",
    ]

    def add_group_section(
        title: str,
        groups: dict[str, dict[str, Any]],
        labels: dict[str, str],
    ) -> None:
        summary.append([])
        summary.append(_excel_row([title]))
        title_row = summary.max_row
        summary.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=9)
        summary.cell(title_row, 1).fill = PatternFill("solid", fgColor=pale_blue)
        summary.cell(title_row, 1).font = Font(bold=True, color=navy)
        summary.append(_excel_row(verdict_headers))
        header_row = summary.max_row
        for cell in summary[header_row]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for key, group in groups.items():
            counts = group["verdict_counts"]
            summary.append(
                _excel_row([
                    labels.get(key, key),
                    group["sample_count"],
                    group["evaluated_count"],
                    group["correct_count"],
                    group["sample_accuracy"],
                    counts["candidate_selected"],
                    counts["all_candidates_incorrect"],
                    counts["no_candidate_correct"],
                    counts["kcs_missing"],
                ])
            )
            summary.cell(summary.max_row, 5).number_format = "0.00%"

    add_group_section("코호트별 통계", insights["cohorts"], QUALITY_COHORT_LABELS)
    add_group_section("점수 구간별 통계", insights["score_bands"], QUALITY_SCORE_BAND_LABELS)

    summary.append([])
    summary.append(_excel_row(["오류 신호"]))
    error_title_row = summary.max_row
    summary.merge_cells(
        start_row=error_title_row, start_column=1, end_row=error_title_row, end_column=3
    )
    summary.cell(error_title_row, 1).fill = PatternFill("solid", fgColor=pale_blue)
    summary.cell(error_title_row, 1).font = Font(bold=True, color=navy)
    summary.append(_excel_row(["신호", "건수", "평가 완료 표본 대비 비율"]))
    error_header_row = summary.max_row
    for cell in summary[error_header_row]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    error_signals = insights["error_signals"]
    for label, key in (
        ("적정 후보가 2순위", "candidate_selected_rank_2"),
        ("적정 후보가 3순위", "candidate_selected_rank_3"),
        ("제시 후보 모두 부적정(오탐)", "all_candidates_incorrect"),
        ("정답 KCS 누락(후보 있음)", "kcs_missing_with_candidates"),
        ("정답 KCS 누락(후보 없음)", "kcs_missing_without_candidates"),
        ("정답 KCS 누락 합계(미탐)", "kcs_missing"),
        ("오류 신호 합계", "total"),
    ):
        rate = (
            error_signals[key] / insights["evaluated_count"]
            if insights["evaluated_count"]
            else None
        )
        summary.append(_excel_row([label, error_signals[key], rate]))
        summary.cell(summary.max_row, 3).number_format = "0.00%"

    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 44
    for column in range(3, 10):
        summary.column_dimensions[get_column_letter(column)].width = 18
    for row in summary.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    detail = workbook.create_sheet("표본평가")
    headers = [
        "표본순서", "코호트", "점수구간", "포스코 조항", "포스코 제목", "포스코 원문",
        "후보1 KCS 코드", "후보1 조항", "후보1 제목", "후보1 score",
        "후보2 KCS 코드", "후보2 조항", "후보2 제목", "후보2 score",
        "후보3 KCS 코드", "후보3 조항", "후보3 제목", "후보3 score",
        "평가결과", "선택후보순위", "누락 KCS 코드", "누락 KCS 조항",
        "평가 메모", "평가일시",
    ]
    detail.append(_excel_row(headers))
    for row in rows:
        candidates = {candidate["rank"]: candidate for candidate in row.get("candidates", [])}
        values: list[Any] = [
            row.get("sample_order"),
            QUALITY_COHORT_LABELS.get(row.get("cohort"), row.get("cohort")),
            QUALITY_SCORE_BAND_LABELS.get(row.get("score_band"), row.get("score_band")),
            row.get("label"),
            row.get("title"),
            row.get("content"),
        ]
        for rank in (1, 2, 3):
            candidate = candidates.get(rank, {})
            values.extend(
                [
                    candidate.get("kcs_code"),
                    candidate.get("kcs_clause"),
                    candidate.get("title"),
                    candidate.get("score"),
                ]
            )
        values.extend(
            [
                QUALITY_VERDICT_LABELS.get(row.get("verdict"), row.get("verdict")),
                row.get("relevant_rank"),
                row.get("expected_kcs_code"),
                row.get("expected_kcs_clause"),
                row.get("quality_note"),
                row.get("evaluated_at"),
            ]
        )
        detail.append(_excel_row(values))

    for cell in detail[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, detail.max_row)}"
    widths = [
        10, 18, 22, 14, 28, 60,
        18, 14, 30, 12, 18, 14, 30, 12, 18, 14, 30, 12,
        24, 14, 20, 18, 40, 24,
    ]
    for index, width in enumerate(widths, start=1):
        detail.column_dimensions[get_column_letter(index)].width = width
    for row in detail.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in (10, 14, 18):
            row[column - 1].number_format = "0.0000"

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
